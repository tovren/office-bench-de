"""docx checker: structure (tables, headings, table of contents, lists, page header) and the text in reading order.

python-docx is imported as the absolute package `docx`; this module's own siblings are imported relatively, so the
module name does not shadow the library."""
import re
import zipfile
from pathlib import Path

from .core import Context, Document, Table
from .registry import SOFT

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
NUMBERED_FORMATS = {"decimal", "decimalZero", "lowerLetter", "upperLetter", "lowerRoman", "upperRoman", "ordinal"}
MANUAL_NUMBER = re.compile(r"^\s*\d{1,2}[.)]\s+\S")
# A heading numbered by hand: '1. Geltungsbereich', '1 Zweck', '2.3 Umfang', '1) …', '§ 3 …', 'IV. …', 'A. …',
# 'Abschnitt 2: …'
HEADING_NUMBER = re.compile(r"^\s*(?:§\s*\d{1,3}|\d{1,2}(?:\.\d{1,2})*\.?|\d{1,2}\)|[IVX]{1,5}\.|[A-H]\.|"
                            r"(?:abschnitt|kapitel|teil|artikel|art\.)\s+\d{1,2}\.?)(?=\s|:|$)", re.I)
TOC_HEADING = re.compile(r"^\s*(?:\d+\.?\s*)?(?:inhaltsverzeichnis|inhalt|inhaltsübersicht|gliederung)\s*:?\s*$", re.I)


def _cell_text(tc) -> str:
    return "\n".join("".join(t.text or "" for t in p.iter(f"{W}t")) for p in tc.iter(f"{W}p"))


def _is_heading(p) -> bool:
    return p.style is not None and bool(re.match(r"(Heading|Überschrift) \d", p.style.name or ""))


def _num_formats(d) -> dict:
    """{(numId, ilvl): numFmt} from numbering.xml."""
    try:
        numbering = d.part.numbering_part.element
    except (KeyError, NotImplementedError, AttributeError):
        return {}
    abstract = {}
    for an in numbering.iter(f"{W}abstractNum"):
        levels = {}
        for lvl in an.iter(f"{W}lvl"):
            fmt = lvl.find(f"{W}numFmt")
            levels[lvl.get(f"{W}ilvl")] = fmt.get(f"{W}val") if fmt is not None else None
        abstract[an.get(f"{W}abstractNumId")] = levels
    out = {}
    for num in numbering.iter(f"{W}num"):
        ref = num.find(f"{W}abstractNumId")
        if ref is not None:
            for ilvl, fmt in abstract.get(ref.get(f"{W}val"), {}).items():
                out[(num.get(f"{W}numId"), ilvl)] = fmt
    return out


def has_numbered_list(d, paragraphs) -> bool:
    """A numbered list is present: list style, Word numbering, or ≥ 2 consecutive '1.' paragraphs."""
    formats = _num_formats(d)
    run = 0
    for p in paragraphs:
        style = (p.style.name if p.style is not None else "") or ""
        if re.match(r"List Number|Listennummer|Liste Nummer", style):
            return True
        ppr = p._p.pPr
        if ppr is not None and ppr.numPr is not None and ppr.numPr.numId is not None:
            ilvl = ppr.numPr.ilvl.val if ppr.numPr.ilvl is not None else 0
            if formats.get((str(ppr.numPr.numId.val), str(ilvl))) in NUMBERED_FORMATS:
                return True
        if _is_heading(p):
            run = 0
            continue
        run = run + 1 if MANUAL_NUMBER.match(p.text or "") else 0
        if run >= 2:
            return True
    return False


def _num_pr(p):
    """(numId, ilvl) of the Word numbering of a paragraph: its own numPr, else the one of its style or a base style;
    None without numbering (numId 0 switches a style's numbering off)."""
    ppr = p._p.pPr
    ilvl = None
    if ppr is not None and ppr.numPr is not None:
        if ppr.numPr.ilvl is not None:
            ilvl = ppr.numPr.ilvl.val
        if ppr.numPr.numId is not None:
            return str(ppr.numPr.numId.val), str(ilvl or 0)
    style, seen = p.style, set()
    while style is not None and id(style.element) not in seen:
        seen.add(id(style.element))
        num = style.element.pPr.numPr if style.element.pPr is not None else None
        if num is not None and num.numId is not None:
            if ilvl is None and num.ilvl is not None:
                ilvl = num.ilvl.val
            return str(num.numId.val), str(ilvl or 0)
        style = style.base_style
    return None


def numbered_headings(d) -> tuple[int, int]:
    """(headings, numbered headings): numbered by hand ('1. Geltungsbereich', '§ 3', 'IV.') or by Word numbering
    (also through the heading style) in a numbered format, not bullets."""
    formats = _num_formats(d)
    n = numbered = 0
    for p in d.paragraphs:
        if not _is_heading(p) or not (p.text or "").strip():
            continue
        n += 1
        num = _num_pr(p)
        if HEADING_NUMBER.match(p.text) or num and num[0] != "0" and formats.get(num) in NUMBERED_FORMATS:
            numbered += 1
    return n, numbered


def header_text(path: Path) -> str:
    """Text of all page headers, read from the XML parts (python-docx only exposes the headers of each section)."""
    with zipfile.ZipFile(path) as z:
        parts = [z.read(n).decode("utf-8", "replace") for n in z.namelist() if re.match(r"word/header\d*\.xml$", n)]
    return "\n".join(" ".join(re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>", x)) for x in parts)


def check_docx(path: Path, ctx: Context) -> tuple[dict, Document]:
    import docx
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph

    d = docx.Document(path)
    doc = Document()
    n_paragraphs = n_tables = 0
    tables_with_content = 0
    for child in d.element.body.iterchildren():
        if child.tag == f"{W}p":
            n_paragraphs += 1
            p = Paragraph(child, d)
            style = (p.style.name if p.style is not None else "") or ""
            where = "title" if style in ("Title", "Subtitle", "Titel", "Untertitel") else f"paragraph {n_paragraphs}"
            doc.paragraphs.append((where, p.text))
        elif child.tag == f"{W}tbl":
            n_tables += 1
            t = DocxTable(child, d)
            rows = []
            for r, row in enumerate(t.rows, 1):
                line, cells, seen = [], row.cells, set()  # keep the cells alive: id() must stay unique
                for c in cells:  # merged cells repeat in python-docx; keep them for the columns, text once
                    text = c.text
                    line.append(text)
                    if id(c._tc) not in seen:
                        seen.add(id(c._tc))
                        doc.paragraphs.append((f"table {n_tables}, row {r}", text))
                rows.append(line)
            doc.tables.append(Table(f"table {n_tables}", rows))
            if any((z or "").strip() for line in rows[1:] for z in line):
                tables_with_content += 1
    headings = sum(1 for p in d.paragraphs if _is_heading(p))
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", "replace")
    toc = bool(re.search(r"<w:instrText[^>]*>\s*TOC\b", xml) or re.search(r'w:instr="\s*TOC\b', xml))
    with ctx.guarded("docx"):
        _task_checks(d, path, ctx, n_tables, tables_with_content, headings, toc)
    return {"headings": headings, "tables": n_tables, "toc": toc}, doc


def _task_checks(d, path: Path, ctx: Context, n_tables: int, tables_with_content: int, headings: int, toc: bool):
    """The docx keys of [checks]."""
    c = ctx.checks
    if "tables_min" in c and tables_with_content < c["tables_min"]:
        if n_tables > tables_with_content:
            ctx.report("tables_min", f"no table with content ({n_tables - tables_with_content} table with header row "
                                     f"only; expected ≥ {c['tables_min']})" if tables_with_content == 0 else
                       f"only {tables_with_content} tables with content (expected ≥ {c['tables_min']})")
        else:
            ctx.report("tables_min", "no table" if n_tables == 0 else
                       f"only {n_tables} tables (expected ≥ {c['tables_min']})")
    if c.get("toc") and not toc:
        if any(TOC_HEADING.match(x.text or "") for x in d.paragraphs):
            # a heading 'Inhaltsverzeichnis' with entries typed by hand (or '[TOC]'): there, but not updatable
            ctx.report("toc", "table of contents typed by hand only (no TOC field, page numbers missing)",
                       label="manual table of contents", severity=SOFT)
        else:
            ctx.report("toc", "no table of contents (TOC field missing)")
    if "headings_min" in c and headings < c["headings_min"]:
        ctx.report("headings_min", f"only {headings} headings (expected ≥ {c['headings_min']})")
    if c.get("numbered_list") and not has_numbered_list(d, d.paragraphs):
        ctx.report("numbered_list", "no numbered list")
    if c.get("numbered_headings"):
        # fewer than two headings: headings_min reports it (sections as plain numbered paragraphs are numbered)
        n, numbered = numbered_headings(d)
        if n >= 2 and numbered < 2:
            ctx.report("numbered_headings", f"sections not numbered: {numbered} of {n} headings with a number "
                                            "(„1.“, „§ 1“ or Word numbering)")
    if "header_contains" in c:
        header = header_text(path)
        if c["header_contains"].lower() not in header.lower():
            ctx.report("header_contains", f"„{c['header_contains']}“ is not in the page header"
                       + (f" (header: „{header.strip()[:40]}“)" if header.strip() else " (no page header)"))
