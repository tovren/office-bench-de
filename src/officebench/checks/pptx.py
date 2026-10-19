"""pptx checker: slides, charts, tables and speaker notes, and the text in reading order; per slide the text and
whether it shows two columns (side_by_side, last_slide_topics).

python-pptx is imported as the absolute package `pptx`; this module's own siblings are imported relatively, so the
module name does not shadow the library."""
import re
from pathlib import Path

from .arithmetic import numbers_in_text
from .charts import chart_info, check_charts
from .core import Context, Document, Slide, Table
from .text import missing_topics, words
from .util import short

# A closing slide at the end ('Vielen Dank', 'Fragen?') does not count as the last slide for last_slide_topics.
CLOSING_SLIDE = re.compile(r"danke|fragen|kontakt|diskussion|aufmerksamkeit|\bende\b|q\s*&\s*a", re.I)
CLOSING_SLIDE_WORDS = 15


def side_by_side(a, b) -> bool:
    """Two boxes (left, top, width, height) side by side: their columns overlap by at most 10 % of the narrower one,
    their rows by at least half of the lower one."""
    (la, ta, wa, ha), (lb, tb, wb, hb) = a, b
    across = min(la + wa, lb + wb) - max(la, lb)
    down = min(ta + ha, tb + hb) - max(ta, tb)
    return across <= 0.1 * min(wa, wb) and down >= 0.5 * min(ha, hb) > 0


def check_pptx(path: Path, ctx: Context) -> tuple[dict, Document]:
    from lxml import etree
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER

    prs = Presentation(path)
    doc = Document()
    charts, tables, tables_with_content = [], 0, 0
    with_notes = 0  # slides whose speaker notes have at least 5 words
    minor = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE}
    titles = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.SUBTITLE}
    doc.slides = []
    for nr, slide in enumerate(prs.slides, 1):
        where = "title slide" if nr == 1 else f"slide {nr}"
        info, texts, boxes = Slide(nr), [], []
        for shape in slide.shapes:
            kind = None
            if shape.is_placeholder:
                try:
                    kind = shape.placeholder_format.type
                except Exception:
                    pass
            if kind in minor:
                continue
            if getattr(shape, "has_chart", False) and shape.has_chart:
                charts.append(chart_info(etree.tostring(shape.chart._chartSpace)))
            elif getattr(shape, "has_table", False) and shape.has_table:
                tables += 1
                rows = []
                for r, row in enumerate(shape.table.rows, 1):
                    line = []
                    for c in row.cells:
                        doc.paragraphs.append((f"{where}, table row {r}", c.text))
                        line.append(c.text)
                    rows.append(line)
                doc.tables.append(Table(f"{where}, table", rows))
                texts += [z for line in rows for z in line if (z or "").strip()]
                filled = [line for line in rows if sum(1 for z in line if (z or "").strip()) >= 2]
                info.two_columns |= len(filled) >= 2
                if any((z or "").strip() for line in rows[1:] for z in line):
                    tables_with_content += 1
            elif shape.has_text_frame and shape.text_frame.text.strip():
                doc.paragraphs.append((where, shape.text_frame.text))
                texts.append(shape.text_frame.text)
                # a column: not a title, at least two words and 15 % of the slide width (no page number or footer)
                if kind not in titles and None not in (shape.left, shape.top, shape.width, shape.height) \
                        and words(shape.text_frame.text) >= 2 and shape.width >= 0.15 * (prs.slide_width or 0):
                    boxes.append((shape.left, shape.top, shape.width, shape.height))
        info.two_columns |= any(side_by_side(a, b) for i, a in enumerate(boxes) for b in boxes[i + 1:])
        info.text = "\n".join(texts)
        if slide.has_notes_slide:
            note = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ""
            doc.paragraphs.append((f"slide {nr}, notes", note))
            with_notes += words(note) >= 5
            info.notes = note
        doc.slides.append(info)
    n = len(prs.slides)
    doc.numbers = numbers_in_text(doc.text) + [x for ch in charts for s in ch.series for x in s.values]
    c = ctx.checks
    with ctx.guarded("pptx"):
        if "slides" in c and not c["slides"][0] <= n <= c["slides"][1]:
            ctx.report("slides", f"{n} slides (expected {c['slides'][0]}–{c['slides'][1]})")
        check_charts(charts, ctx)
        if "notes_min" in c and with_notes < c["notes_min"]:
            ctx.report("notes_min", f"only {with_notes} slides with speaker notes of at least 5 words "
                                    f"(expected ≥ {c['notes_min']})")
        if "tables_min" in c and tables_with_content < c["tables_min"]:
            ctx.report("tables_min", "no table" if tables == 0 else
                       f"only {tables_with_content} tables with content (expected ≥ {c['tables_min']})")
    if "side_by_side" in c:
        with ctx.guarded("side_by_side"):
            check_side_by_side(doc.slides, ctx)
    if "last_slide_topics" in c:
        with ctx.guarded("last_slide_topics"):
            check_last_slide(doc.slides, ctx)
    if "slides_added" in c:
        with ctx.guarded("slides_added"):
            check_slides_added(n, ctx)
    return {"slides": n, "charts": len(charts), "tables": tables, "notes": with_notes}, doc


def check_side_by_side(slides: list, ctx: Context) -> None:
    """side_by_side = { name = pattern, … }: one slide shows two columns (text shapes side by side or a table) and its
    text has every topic."""
    topics = ctx.checks["side_by_side"]
    if any(s.two_columns and not missing_topics(s.text, topics) for s in slides):
        return
    names = " and ".join(f"„{n}“" for n in topics)
    with_topics = [s.number for s in slides if not missing_topics(s.text, topics)]
    extra = (f"; the topics are on slide {', '.join(map(str, with_topics[:3]))}, but not in two columns or a table"
             if with_topics else "; no slide names both")
    ctx.report("side_by_side", f"no slide sets {names} side by side in two columns or a table{extra}")


def last_content_slide(slides: list):
    """The last slide, or the one before it if the last is a short closing slide ('Vielen Dank', 'Fragen?')."""
    if len(slides) >= 2 and words(slides[-1].text) <= CLOSING_SLIDE_WORDS and CLOSING_SLIDE.search(slides[-1].text):
        return slides[-2]
    return slides[-1] if slides else None


def check_last_slide(slides: list, ctx: Context) -> None:
    """last_slide_topics: every topic on the last slide (slide text, not the notes; a closing slide after it is
    skipped)."""
    slide = last_content_slide(slides)
    if slide is None:
        return
    missing = missing_topics(slide.text, ctx.checks["last_slide_topics"])
    if missing:
        start = short(" ".join(slide.text.split()), 50)
        ctx.report("last_slide_topics", f"missing on the last slide (slide {slide.number}: „{start}“): "
                                        f"{', '.join(missing)}")


def check_slides_added(n: int, ctx: Context) -> None:
    """slides_added = 1 or [from, to]: slides added compared with the file of the previous turn."""
    from pptx import Presentation

    target = ctx.checks["slides_added"]
    low, high = (target, target) if isinstance(target, int) else target
    if ctx.previous is None or Path(ctx.previous).suffix.lower() != ".pptx":
        ctx.not_checkable("slides_added", "no previous PowerPoint file to compare with")
        return
    try:
        old = len(Presentation(ctx.previous).slides)
    except Exception as e:  # noqa: BLE001 - a broken previous file only makes the check impossible
        ctx.not_checkable("slides_added", f"previous file not readable ({type(e).__name__})")
        return
    if not low <= n - old <= high:
        expected = f"+{low}" if low == high else f"+{low} to +{high}"
        ctx.report("slides_added", f"{n} slides after {old} in the previous file ({n - old:+d}, expected {expected})")
