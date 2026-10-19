"""retained (follow-up tasks): share of the text units of the previous file that the new file still contains.

Units: docx paragraphs and table cells, pptx paragraphs of the shapes, table cells and notes (each split into lines),
xlsx cells with text or constants (no formulas: references shift when a column is inserted). Normalised (lower case,
whitespace collapsed, punctuation at the edges removed); units under 3 characters and numbers with up to 2 digits
are ignored, each distinct unit counts once. A unit is kept if it occurs again or a unit of the new file is at least
90 % similar (difflib)."""
import bisect
import difflib
import math
import re
from datetime import date, datetime
from pathlib import Path

from .arithmetic import format_number
from .core import Context, Document
from .registry import SOFT
from .util import normalize, short

SIMILAR = 0.9
EDGES = re.compile(r"^[\W_]+|[\W_]+$")


def _cell_text(v) -> str:
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        return format_number(float(v))  # 3200 and 3200.0 alike
    if isinstance(v, (datetime, date)):
        return v.strftime("%d.%m.%Y")
    return str(v)


def texts(path: Path, doc: Document | None = None) -> list[str]:
    """Raw text units of a file (doc: the already read document of a docx/pptx)."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path)  # stored values are enough: formulas are skipped
        return [_cell_text(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
                if c.value is not None and c.data_type != "f"]
    if doc is None:
        from .docx import check_docx
        from .pptx import check_pptx

        checker = {".docx": check_docx, ".pptx": check_pptx}.get(suffix)
        if checker is None:
            raise ValueError(f"unknown file type {path.suffix}")
        _, doc = checker(path, Context())  # a throwaway context: only the text is wanted
    return [line for _, text in doc.paragraphs for line in text.splitlines()]


def units(raw: list[str]) -> dict[str, str]:
    """{normalised unit: first original} without units under 3 characters and numbers with up to 2 digits."""
    out = {}
    for t in raw:
        n = EDGES.sub("", " ".join(normalize(t).lower().split()))
        if len(n) >= 3 and not re.fullmatch(r"\d{1,2}", n):
            out.setdefault(n, " ".join(t.split()))
    return out


def compare(old: list[str], new: list[str]) -> tuple[float | None, list[str], int]:
    """(share of the old units kept, lost units as written, number of old units); share None without old units."""
    a, n = units(old), set(units(new))
    by_length = sorted(n, key=len)
    lengths = [len(x) for x in by_length]
    sm = difflib.SequenceMatcher(autojunk=False)
    lost = []
    for normed, original in a.items():
        if normed in n:
            continue
        # ratio = 2·M / (la + lb) ≥ 0.9 needs 9/11 · la ≤ lb ≤ 11/9 · la
        lo = bisect.bisect_left(lengths, math.ceil(len(normed) * 9 / 11))
        hi = bisect.bisect_right(lengths, math.floor(len(normed) * 11 / 9))
        sm.set_seq2(normed)  # cached side: one old unit against many new ones
        for cand in by_length[lo:hi]:
            sm.set_seq1(cand)
            if sm.real_quick_ratio() >= SIMILAR and sm.quick_ratio() >= SIMILAR and sm.ratio() >= SIMILAR:
                break
        else:
            lost.append(original)
    return ((len(a) - len(lost)) / len(a) if a else None), lost, len(a)


def check_retained(path: Path, doc: Document | None, ctx: Context) -> dict:
    """retained = 0.8 (hard lower bound) or { hard = 0.5, soft = 0.8 }; returns the metric 'retained'."""
    p = ctx.checks["retained"]
    levels = p if isinstance(p, dict) else {"hard": p}
    if ctx.previous is None:
        ctx.not_checkable("retained", "no previous file to compare with")
        return {}
    try:
        old = texts(Path(ctx.previous))
    except Exception as e:  # noqa: BLE001 - a missing or broken previous file only makes the check impossible
        ctx.not_checkable("retained", f"previous file not readable ({type(e).__name__}: {str(e)[:80]})")
        return {}
    share, lost, n = compare(old, texts(path, doc))
    if share is None:
        ctx.not_checkable("retained", "previous file has no text")
        return {}
    for level in ("hard", "soft"):
        if level in levels and share < levels[level]:
            examples = ", ".join(f"„{short(x, 40)}“" for x in lost[:2])
            ctx.report("retained", f"only {100 * share:.0f} % of the previous file's content retained ({len(lost)} of "
                                   f"{n} units missing, expected ≥ {100 * levels[level]:.0f} %), e.g. {examples}",
                       severity=None if level == "hard" else SOFT)
            break
    return {"retained": f"{100 * share:.0f} %"}
