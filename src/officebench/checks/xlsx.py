"""xlsx checks. Workbooks are recalculated with LibreOffice headless first (xlsx -> xlsx into a temp folder, then
openpyxl data_only=True): files written by openpyxl carry no cached formula results, so only a recalculation makes
results such as #VALUE! visible and lets the checks redo calculations."""

import ast
import itertools
import os
import re
import shutil
import signal
import subprocess
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from .arithmetic import evaluate, format_number, number_values
from .charts import chart_info, check_charts
from .core import Context, Document, Table
from .dates import dates_in
from .registry import HARD, SOFT
from .util import MONTH_NAMES

_MACOS_SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"


def _find_soffice() -> str:
    """soffice on PATH, else the macOS app bundle (which is not on PATH by default), else the Linux default path."""
    if found := shutil.which("soffice"):
        return found
    return _MACOS_SOFFICE if Path(_MACOS_SOFFICE).exists() else "/usr/bin/soffice"


SOFFICE = _find_soffice()
ERROR_VALUE = re.compile(r"^(#VALUE!|#REF!|#NAME\?|#DIV/0!|#N/A|#NUM!|#NULL!|#WERT!|#BEZUG!|#NV|Err:\d+)")
NUMBER_TEXT = re.compile(
    r"^\s*[-+−]?\s*(?:€|EUR|\$)?\s*(?:\d{1,3}(?:[.\s'’ ]\d{3})+|\d+)(?:[.,]\d+)?\s*"
    r"(?:€|EUR|Euro|%|\$|USD|CHF|h|Std\.?|Stk\.?)?\s*$", re.I)
ID_COLUMN = re.compile(r"\b(pos|pos\.|position|nr|nr\.|nummer|id|jahr|plz|telefon|tel|version|kw|kalenderwoche)\b",
                       re.I)
NUMBER_HEADER = re.compile(r"€|\beur\b|euro|preis|betrag|kosten|summe|saldo|einnahm|ausgab|menge|anzahl|stunden|"
                           r"satz\b|umsatz|budget|gesamt|netto|brutto|\bstk\b|%|prozent", re.I)
SUM = re.compile(r"\b(SUM|SUMME|SUBTOTAL|TEILERGEBNIS)\s*\(", re.I)
OTHER_FUNCTION = re.compile(r"^=\s*(AVERAGE|MITTELWERT|MIN|MAX|COUNTA?|ANZAHL2?|MEDIAN)\s*\(", re.I)
RANGE = re.compile(r"(?:(?P<blatt>'[^']+'|[A-Za-z0-9_.ÄÖÜäöüß]+)!)?\$?(?P<c1>[A-Z]{1,3})\$?(?P<r1>\d+):"
                   r"\$?(?P<c2>[A-Z]{1,3})\$?(?P<r2>\d+)")
# Labels of total rows in sheets; not the same list as the total rows of Word/slide tables in arithmetic.py, merging
# them would change results.
TOTAL_LABEL = re.compile(r"^(summe|gesamt|gesamtsumme|total|netto|netto-summe|nettosumme|brutto|zwischensumme)\b",
                         re.I)
FORMULA_TEXT = re.compile(r"^[\"'„“”]?\s*=\s*(?:[A-Za-z][A-Za-z0-9.]*\(|\$?[A-Z]{1,3}\$?\d+)")
MONTHS = [(1, r"jan(?:uar|\.)?|jän(?:ner)?"), (2, r"feb(?:ruar|\.)?"), (3, r"mär(?:z|\.)?|maerz|mrz|mar"),
          (4, r"apr(?:il|\.)?"), (5, r"mai"), (6, r"jun(?:i|\.)?"), (7, r"jul(?:i|\.)?"), (8, r"aug(?:ust|\.)?"),
          (9, r"sep(?:t(?:ember|\.)?|\.)?"), (10, r"okt(?:ober|\.)?|oct(?:ober)?"), (11, r"nov(?:ember|\.)?"),
          (12, r"dez(?:ember|\.)?|dec(?:ember)?")]
MONTH_RES = [(n, re.compile(rf"(?<![A-Za-zÄÖÜäöüß])(?:{p})(?![A-Za-zÄÖÜäöüß])", re.I)) for n, p in MONTHS]
MONTH_NUMBER = re.compile(r"\b(?:(0?[1-9]|1[0-2])[./-]20\d\d|20\d\d-(0[1-9]|1[0-2]))\b")
SHORT_MONTH = {n: MONTH_NAMES[n][:3] for n in range(1, 13)}
# A chart reference: 'Blatt 1'!$B$2:$B$7, Tabelle1!B2 (one area; several areas are split at commas first).
REFERENCE = re.compile(r"^(?:'(?P<q>(?:[^']|'')+)'|(?P<n>[^'!]+))!\$?(?P<c1>[A-Z]{1,3})\$?(?P<r1>\d+)"
                       r"(?::\$?(?P<c2>[A-Z]{1,3})\$?(?P<r2>\d+))?$")


def col_letter(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def col_number(s: str) -> int:
    n = 0
    for ch in s:
        n = n * 26 + ord(ch) - 64
    return n


def runs(numbers) -> tuple:
    """Consecutive runs: {2, 3, 4, 7} -> ((2, 4), (7, 7))."""
    numbers = sorted(set(numbers))
    parts, start, prev = [], numbers[0], numbers[0]
    for z in numbers[1:]:
        if z != prev + 1:
            parts.append((start, prev))
            start = z
        prev = z
    parts.append((start, prev))
    return tuple(parts)


def col_range(sheet: str, col: int, rows) -> str:
    return cell_ranges(sheet, [(col, r) for r in rows])


def cell_ranges(sheet: str, cells, max_parts: int = 6) -> str:
    """Compact A1 ranges for (column, row) cells; neighbouring columns with the same rows become rectangles, so a
    defect names a whole block instead of listing every cell."""
    per_col = defaultdict(set)
    for col, r in cells:
        per_col[col].add(r)
    cols, parts, i = sorted(per_col), [], 0
    while i < len(cols):
        rows, j = runs(per_col[cols[i]]), i
        while j + 1 < len(cols) and cols[j + 1] == cols[j] + 1 and runs(per_col[cols[j + 1]]) == rows:
            j += 1
        c0, c1 = col_letter(cols[i]), col_letter(cols[j])
        parts += [f"{c0}{a}" if (c0, a) == (c1, b) else f"{c0}{a}:{c1}{b}" for a, b in rows]
        i = j + 1
    rest = f" and {len(parts) - max_parts} more" if len(parts) > max_parts else ""
    return f"{sheet}!" + ", ".join(parts[:max_parts]) + rest


def _cells(sheet: str, cells) -> str:
    return cell_ranges(sheet, [(z.column, z.row) for z in cells])


def empty(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def is_formula(c) -> bool:
    """A real formula cell (cell type), not text that merely starts with '='."""
    return c.data_type == "f"


def percent_format(fmt) -> bool:
    """A percent number format (0.0%); a quoted or escaped % ('0.00" %"') is only text and does not multiply."""
    return "%" in re.sub(r'"[^"]*"|\\.', "", fmt or "")


def is_number_format(fmt) -> bool:
    return bool(fmt) and fmt not in ("General", "@") and bool(re.search(r"[0#]", fmt))


def number(v) -> float | None:
    """Number of a cell value; numbers stored as text ('1.500 €') count too."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and NUMBER_TEXT.match(v):
        raw = re.sub(r"[^\d.,'’-]", "", v.replace("−", "-"))
        neg = raw.startswith("-")
        values = number_values(raw.lstrip("-").replace("’", "'"))
        if values:
            return -values[0] if neg else values[0]
    return None


def number_or_date(v) -> float | None:
    """Like number(), but dates count as Excel serial numbers (days since 30.12.1899), so 'Ende = Start + Dauer'
    can be recalculated."""
    if isinstance(v, datetime):
        return (v - datetime(1899, 12, 30)).total_seconds() / 86400
    if isinstance(v, date):
        return float((v - date(1899, 12, 30)).days)
    return number(v)


def months_in(value, lo: int = 1, hi: int = 12) -> set:
    if isinstance(value, (datetime, date)):
        return {value.month} if lo <= value.month <= hi else set()
    if not isinstance(value, str) or value.startswith("="):
        return set()
    out = {n for n, rx in MONTH_RES if rx.search(value)}
    out |= {int(a or b) for a, b in MONTH_NUMBER.findall(value)}
    return {n for n in out if lo <= n <= hi}


def _closes_block(formula: str, row: int, start: int) -> bool:
    """A vertical SUM over the rows of its block, ending above `row` (a total, not a running sum)."""
    if not SUM.search(formula):
        return False
    return any(int(m["r1"]) != int(m["r2"]) and int(m["r2"]) < row and int(m["r1"]) <= start + 1
               for m in RANGE.finditer(formula) if not m["blatt"])


def display(v) -> str:
    if isinstance(v, (datetime, date)):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float):
        return format_number(v)
    return "" if v is None else str(v)


class Sheet:
    """Header row, data rows and total rows of one sheet (header = first non-empty row).

    A total row has a label like 'Summe'/'Gesamt' or a vertical SUM that ends above the row and starts at the first
    row of its block (after the header or the previous total row). A running total like =SUM(D$2:D3) in row 3 is
    no total row, otherwise every row of a cumulative column would count as a total."""

    def __init__(self, ws, wsv=None):
        self.ws, self.wsv, self.name = ws, wsv, ws.title
        self.width = ws.max_column
        self.header_row, self.headers, self.data, self.totals = None, {}, [], []
        for r in range(1, ws.max_row + 1):
            cells = [ws.cell(r, c) for c in range(1, self.width + 1)]
            if all(empty(z.value) for z in cells):
                continue
            if self.header_row is None:
                self.header_row = r
                self.headers = {z.column: str(z.value).strip() for z in cells if not empty(z.value)}
                continue
            first = next((z.value for z in cells if not empty(z.value)), "")
            start = self.totals[-1] if self.totals else self.header_row
            if any(is_formula(z) and _closes_block(str(z.value), r, start) for z in cells) or (
                    isinstance(first, str) and TOTAL_LABEL.match(first.strip())):
                self.totals.append(r)
            else:
                self.data.append(r)

    def header(self, col: int) -> str:
        return self.headers.get(col, col_letter(col))

    def formulas(self):
        for row in self.ws.iter_rows():
            for c in row:
                if is_formula(c):
                    yield c

    def value(self, r: int, c: int):
        """Recalculated value if available, else the stored value (None for formulas)."""
        cell = self.ws.cell(r, c)
        if self.wsv is not None:
            return self.wsv.cell(r, c).value
        return None if is_formula(cell) else cell.value

    def label(self, r: int) -> tuple[int, str] | None:
        """(column, text) of the first non-empty cell of a row if it is text."""
        for c in range(1, self.width + 1):
            v = self.ws.cell(r, c).value
            if not empty(v):
                return (c, v.strip()) if isinstance(v, str) and not is_formula(self.ws.cell(r, c)) else None
        return None


def _resolver(sheets: list):
    """Resolver for chart references: formula -> ([cell values], header above / label left of the range) or None,
    with the recalculated values (Sheet.value); several areas ('(A!$B$2,A!$B$4)') are joined."""
    by_name = {b.name.lower(): b for b in sheets}

    def resolve(formula: str):
        values, head = [], ""
        for part in re.split(r",(?=(?:[^']*'[^']*')*[^']*$)", formula.strip().strip("()")):
            m = REFERENCE.match(part.strip())
            b = by_name.get(((m["q"] or "").replace("''", "'") or m["n"]).lower()) if m else None
            if b is None:
                return None
            c1, r1 = col_number(m["c1"]), int(m["r1"])
            c2, r2 = (col_number(m["c2"]), int(m["r2"])) if m["c2"] else (c1, r1)
            for r in range(r1, min(r2, b.ws.max_row) + 1):  # cells beyond the sheet are empty (and not created)
                values += [b.value(r, c) for c in range(c1, min(c2, b.ws.max_column) + 1)]
            if not head:
                v = b.value(r1 - 1, c1) if c1 == c2 and r1 > 1 else b.value(r1, c1 - 1) if r1 == r2 and c1 > 1 else None
                head = v.strip() if isinstance(v, str) else ""
        return values, head
    return resolve


def require_soffice(soffice: str = SOFFICE, purpose: str = "") -> None:
    if not Path(soffice).exists() and not shutil.which(soffice):
        raise RuntimeError(f"LibreOffice not found ({soffice})" + (f"; {purpose}" if purpose else ""))


def run_soffice(args: list, profile: Path, timeout: float, soffice: str = SOFFICE) -> None:
    """LibreOffice headless with `args` and the profile folder `profile` (fresh_profile), in its own process group;
    on timeout kill the whole group (soffice.bin, not only oosplash), else a hanging soffice.bin keeps the profile
    locked and blocks the retry."""
    command = [soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--norestore", "--nologo",
               "--nodefault", *map(str, args)]
    proc = subprocess.Popen(command, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
    try:
        proc.wait(timeout=timeout)
    except subprocess.TimeoutExpired:
        try:
            os.killpg(proc.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        proc.wait()


def fresh_profile(folder: Path) -> Path:
    """A fresh LibreOffice profile that recalculates formulas when loading xlsx (OOXMLRecalcMode 0 = always); the
    user's own profile would keep cached results and must not be touched or locked."""
    (folder / "user").mkdir(parents=True, exist_ok=True)
    (folder / "user" / "registrymodifications.xcu").write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
        'xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
        '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse">'
        '<value>0</value></prop></item>\n</oor:items>\n')
    return folder


def recalculate(paths: list, tmp: Path, soffice: str = SOFFICE, timeout: float = 900) -> dict:
    """Recalculate all xlsx with one LibreOffice call (per 300 files, since every start costs seconds); returns
    {original: recalculated or None}. Files missing afterwards (LibreOffice stopped or timed out) get a second try
    with a fresh profile in smaller batches. Profiles and copies live below `tmp` (a temp dir of the caller)."""
    if not paths:
        return {}
    require_soffice(soffice, "xlsx files cannot be recalculated")
    src, out = tmp / "in", tmp / "out"
    for d in (src, out):
        d.mkdir(parents=True, exist_ok=True)
    mapping, copies = {}, []
    for i, p in enumerate(dict.fromkeys(paths)):
        copy = src / f"d{i:05d}.xlsx"
        shutil.copyfile(p, copy)
        mapping[p] = out / copy.name
        copies.append(copy)

    def convert(files: list, per_call: int, profile: Path):
        for start in range(0, len(files), per_call):
            run_soffice(["--convert-to", "xlsx", "--outdir", out, *files[start:start + per_call]], profile, timeout,
                        soffice)

    convert(copies, 300, fresh_profile(tmp / "profile"))
    missing = [k for k in copies if not (out / k.name).exists()]
    if missing:
        convert(missing, 20, fresh_profile(tmp / "profile2"))
    return {p: (q if q.exists() else None) for p, q in mapping.items()}


# ------------------------------------------------------------------------------------ general checks --

def general_checks(b: Sheet, ctx: Context) -> None:
    """formula_errors, formula_empty, formula_as_text, number_as_text, text_in_number_column, empty_number_column,
    zero_sum, sum_range, invalid_date, date_as_text, share_over_100 for one sheet (each guarded: a bug in one check
    only gives a soft defect)."""
    with ctx.guarded("formula_errors"):
        _formula_errors(b, ctx)
    with ctx.guarded("formula_as_text"):
        as_text = [c for row in b.ws.iter_rows() for c in row if c.data_type in ("s", "inlineStr", "str")
                   and isinstance(c.value, str) and FORMULA_TEXT.match(c.value)]
        if as_text and ctx.enabled("formula_as_text"):
            ctx.report("formula_as_text", f"formula stored as text (not calculated): "
                                          f"{_cells(b.name, as_text)} "
                                          f"(e.g. „{as_text[0].value[:30]}“)")
    with ctx.guarded("number_as_text"):
        _column_contents(b, ctx)
    with ctx.guarded("zero_sum"):
        _zero_sum(b, ctx)
    if ctx.enabled("sum_range"):
        with ctx.guarded("sum_range"):
            sum_range(b, ctx)
    with ctx.guarded("invalid_date"):
        _date_text(b, ctx)
    if ctx.enabled("share_over_100"):
        with ctx.guarded("share_over_100"):
            _share_over_100(b, ctx)


def _formula_errors(b: Sheet, ctx: Context) -> None:
    ws, wsv = b.ws, b.wsv
    if wsv is None:
        return
    errors, no_result = defaultdict(list), defaultdict(list)
    for c in b.formulas():
        v = wsv[c.coordinate].value
        if isinstance(v, str) and ERROR_VALUE.match(v):
            errors[(ERROR_VALUE.match(v).group(1), c.column)].append(c.row)
        elif empty(v):
            no_result[c.column].append(c.row)
    for row in wsv.iter_rows():  # error values outside formula cells
        for c in row:
            if isinstance(c.value, str) and ERROR_VALUE.match(c.value) and not is_formula(ws[c.coordinate]):
                errors[(ERROR_VALUE.match(c.value).group(1), c.column)].append(c.row)
    per_error = defaultdict(list)
    for (err, col), rows in errors.items():
        per_error[err] += [(col, r) for r in rows]
    if ctx.enabled("formula_errors"):
        for err, cells in sorted(per_error.items()):
            ctx.report("formula_errors", f"formula error {err} in {cell_ranges(b.name, cells)}")
    if no_result and ctx.enabled("formula_empty"):
        cells = [(col, r) for col, rows in no_result.items() for r in rows]
        ctx.report("formula_empty", f"formula without result in {cell_ranges(b.name, cells)}")


def _column_contents(b: Sheet, ctx: Context) -> None:
    """number_as_text, text_in_number_column, empty_number_column over the data rows."""
    ws = b.ws
    as_text, text_in_number, text_heads = [], [], []
    for col in range(1, b.width + 1):
        cells = [ws.cell(r, col) for r in b.data]
        if not cells:
            continue
        head = b.header(col)
        before = len(text_in_number)
        formatted = sum(1 for z in cells if is_number_format(z.number_format))
        number_col_fmt = formatted >= max(1, len(cells) / 2)
        for z in cells:
            if isinstance(z.value, str) and not is_formula(z):
                if NUMBER_TEXT.match(z.value):
                    if not ID_COLUMN.search(head):
                        as_text.append(z)
                elif number_col_fmt and z.value.strip() and not FORMULA_TEXT.match(z.value):
                    text_in_number.append(z)
        if len(text_in_number) > before:
            text_heads.append(f"„{head}“")
        number_col = number_col_fmt or bool(NUMBER_HEADER.search(head))
        if number_col and not any(is_formula(z) for z in cells) and all(empty(z.value) for z in cells) \
                and ctx.enabled("empty_number_column"):
            ctx.report("empty_number_column",
                       f"number column „{head}“ without values ({col_range(b.name, col, b.data)})")
    # every number typed into the data rows is 0: an empty template, not a filled table (one zero column is fine)
    inputs = [v for col in range(1, b.width + 1) if not ID_COLUMN.search(b.header(col))
              for r in b.data if not is_formula(z := ws.cell(r, col)) and (v := number(z.value)) is not None]
    if len(b.data) >= 2 and len(inputs) >= 2 and all(v == 0 for v in inputs) and ctx.enabled("empty_number_column"):
        ctx.report("empty_number_column", f"no values: all {len(inputs)} numbers in the data rows of {b.name} are 0",
                   label="only zeros", severity=HARD)
    if as_text and ctx.enabled("number_as_text"):
        ctx.report("number_as_text", f"numbers stored as text in {_cells(b.name, as_text)} "
                                     f"(e.g. „{as_text[0].value.strip()}“)")
    if text_in_number and ctx.enabled("text_in_number_column"):
        heads = ", ".join(text_heads[:3]) + (" …" if len(text_heads) > 3 else "")
        ctx.report("text_in_number_column", f"text in number column {heads}: "
                                            f"{_cells(b.name, text_in_number)} "
                                            f"(e.g. „{text_in_number[0].value.strip()[:30]}“)")


def _zero_sum(b: Sheet, ctx: Context) -> None:
    """SUM gives 0 because the range holds numbers or formulas stored as text (a range of zeros is fine)."""
    ws, wsv = b.ws, b.wsv
    if wsv is None or not ctx.enabled("zero_sum"):
        return
    for r in b.totals:
        for c in ws[r]:
            if not (is_formula(c) and SUM.search(str(c.value))):
                continue
            v = wsv[c.coordinate].value
            if not (isinstance(v, (int, float)) and not isinstance(v, bool) and v == 0):
                continue
            for g in RANGE.finditer(str(c.value)):
                if g["blatt"]:
                    continue
                as_text = [ws.cell(rr, cc).value for rr in range(int(g["r1"]), int(g["r2"]) + 1)
                           for cc in range(col_number(g["c1"]), col_number(g["c2"]) + 1)
                           if isinstance(ws.cell(rr, cc).value, str) and not is_formula(ws.cell(rr, cc))
                           and (NUMBER_TEXT.match(ws.cell(rr, cc).value) or FORMULA_TEXT.match(ws.cell(rr, cc).value))]
                if as_text:
                    ctx.report("zero_sum", f"sum = 0 in {b.name}!{c.coordinate} because {g.group(0)} holds numbers or "
                                           f"formulas as text (e.g. „{as_text[0][:20]}“)")
                    break


def _single_range(formula: str):
    """(c1, r1, c2, r2) of =SUM(X1:Y2) with one range on the same sheet, else None."""
    if not re.match(r"^=\s*(SUM|SUMME)\s*\(", formula, re.I):
        return None
    ranges = list(RANGE.finditer(formula))
    if len(ranges) != 1 or ranges[0]["blatt"] or re.search(r"[+\-*/]", formula.split(")", 1)[-1]):
        return None
    g = ranges[0]
    return col_number(g["c1"]), int(g["r1"]), col_number(g["c2"]), int(g["r2"])


# A cell that is only a date written as text: 31.11.2026, 31.11.26, 31.11., 2026-11-31 (also with a time).
DATE_TEXT = re.compile(r"\s*(?:(?P<t>\d{1,2})\.\s?(?P<m>\d{1,2})\.(?:\s?(?P<j>\d{4}|\d{2}))?"
                       r"|(?P<j2>\d{4})-(?P<m2>\d{1,2})-(?P<t2>\d{1,2})(?:[T ]\d{1,2}:\d{2}(?::\d{2})?)?)\s*")


def _date_from_text(text: str) -> tuple[bool, bool] | None:
    """(valid, four-digit year) of a text that is only a date; None if it is no date (month 13, day 32)."""
    m = DATE_TEXT.fullmatch(text)
    if not m:
        return None
    d, mo, y = (m["t"], m["m"], m["j"]) if m["t"] else (m["t2"], m["m2"], m["j2"])
    d, mo, year = int(d), int(mo), int(y) if y else 2024  # 2024: 29.02. without a year stays valid
    if not (1 <= mo <= 12 and 1 <= d <= 31):
        return None
    try:
        date(year if year >= 100 else 2000 + year, mo, d)
    except ValueError:
        return False, bool(y) and len(y) == 4
    return True, bool(y) and len(y) == 4


def _date_text(b: Sheet, ctx: Context) -> None:
    """invalid_date: a text cell with only a date that does not exist (31.11.2026; with a four-digit year in any
    column, else in a date column). date_as_text: a valid date as text in a date column, unless the task's
    column_format checks the column for the date format already (one defect is enough)."""
    invalid, as_text = [], []
    formats = [m for m, kind in (ctx.get("column_format") or {}).items() if kind == "date"]
    for c in range(1, b.width + 1):
        head = b.headers.get(c, "")
        texts = [(z, x) for r in b.data if not is_formula(z := b.ws.cell(r, c)) and isinstance(z.value, str)
                 and (x := _date_from_text(z.value)) is not None]
        if not texts:
            continue
        filled = sum(1 for r in b.data if not empty(b.ws.cell(r, c).value))
        date_col = _is_date_column(b, c) or len(texts) >= max(2, filled / 2)
        for z, (valid, with_year) in texts:
            if not valid and (date_col or with_year):
                invalid.append(z)
            elif valid and date_col and not any(re.search(m, head, re.I) for m in formats):
                as_text.append(z)
    if invalid and ctx.enabled("invalid_date"):
        ctx.report("invalid_date", f"date that does not exist: {_cells(b.name, invalid)} "
                                   f"(e.g. „{invalid[0].value.strip()}“)")
    if as_text and ctx.enabled("date_as_text"):
        ctx.report("date_as_text", f"date stored as text in {_cells(b.name, as_text)} "
                                   f"(e.g. „{as_text[0].value.strip()}“)")


# Percent columns that hold a share (at most 100 %) or a percentage that may exceed 100 % a bit (Auslastung 110 %),
# but not by a factor of ten; changes and target attainment (Veränderung +150 %, Zielerreichung 105 %) are left out.
SHARE_HEADER = re.compile(r"anteil|quote|verteilung", re.I)
PERCENT_HEADER = re.compile(r"auslastung|belegung|prozent|%", re.I)
# The old checker had „\bdiff\b|\bplan\b|\bsoll\b|\bist\b|\bvs\b“ here with literal backspace characters (0x08)
# instead of word boundaries, so those alternatives never matched; left out to keep the verdicts identical (with them
# „Anteil verplant“ or „Quote Liste“ would no longer be checked).
NO_SHARE_HEADER = re.compile(r"veränder|änderung|abweich|wachstum|zuwachs|steiger|delta|differenz|vorjahr|"
                             r"erfüll|erreich|ziel|index|\+/-|±", re.I)


def _share_over_100(b: Sheet, ctx: Context) -> None:
    """share_over_100: percent-formatted values of a share column (Anteil, Quote) over 100 %, of other percent
    columns (Auslastung, „in %“) over 1.000 %: mostly a share multiplied by 100 and shown in the percent format."""
    found = []
    for c, head in sorted(b.headers.items()):
        if NO_SHARE_HEADER.search(head):
            continue
        limit = 1.0 if SHARE_HEADER.search(head) else 10.0 if PERCENT_HEADER.search(head) else None
        if limit is None:
            continue
        too_high = [(r, v) for r in b.data if percent_format(b.ws.cell(r, c).number_format)
                    and isinstance(v := b.value(r, c), (int, float)) and not isinstance(v, bool) and v > limit + 0.005]
        if too_high:
            r, v = too_high[0]
            found.append(f"„{head}“ {col_range(b.name, c, [x for x, _ in too_high])} (e.g. {col_letter(c)}{r} = "
                         f"{format_number(100 * v)} %)")
    if found:
        ctx.report("share_over_100", "share in percent format over 100 %: " + "; ".join(found[:3])
                   + (f" and {len(found) - 3} more" if len(found) > 3 else ""))


# Headers of columns that hold a key or a quantity, not an amount (Fläche, Anteil in %, Mitarbeitende, Stunden …).
KEY_HEADER = re.compile(r"fläche|m²|qm|anteil|%|prozent|quote|anzahl|köpfe|mitarbeit|personen|menge|stück|stunden|"
                        r"schlüssel|\bma\b", re.I)


def _format_kind(fmt: str) -> str:
    fmt = fmt or ""
    if percent_format(fmt):
        return "percent"
    if "€" in fmt or "EUR" in fmt.upper():
        return "euro"
    return "number"


def _same_kind(b: Sheet, r: int, c: int, summed) -> bool:
    """Is column c (left of a SUM range in row r) of the same kind as the summed columns: no key/quantity header and
    the number format (euro, percent, plain) of the summed cells?"""
    if KEY_HEADER.search(b.header(c)):
        return False
    kinds = {_format_kind(b.ws.cell(r, x).number_format) for x in summed}
    return _format_kind(b.ws.cell(r, c).number_format) in kinds


def sum_range(b: Sheet, ctx: Context) -> None:
    """SUM in a total row (or total column) leaves out data rows above (columns left of it), its value differs from
    the sum of the data cells (numbers as text, subtotals counted twice), or a 'Summe' row computes AVERAGE etc."""
    ws = b.ws
    gaps, deviations, others = [], [], []
    total_cols = {c for c, head in b.headers.items() if TOTAL_LABEL.match(head)}
    for r in b.totals:
        lab = b.label(r)
        previous = max([x for x in b.totals if x < r] + [b.header_row or 0])
        block = [x for x in b.data if previous < x < r]
        for cell in ws[r]:
            if not is_formula(cell):
                continue
            formula = str(cell.value)
            if lab and TOTAL_LABEL.match(lab[1]) and (m := OTHER_FUNCTION.match(formula)):
                others.append(f"{b.name}!{cell.coordinate} ={m.group(1).upper()}(…) in the row „{lab[1]}“")
                continue
            g = _single_range(formula)
            if not g or g[0] != g[2] or g[0] != cell.column or g[3] >= r:
                continue
            col = cell.column
            with_number = [x for x in block if number(b.value(x, col)) is not None]
            missing = [x for x in with_number if not g[1] <= x <= g[3]]
            if missing:
                gaps.append(f"{b.name}!{cell.coordinate} ={formula[1:40]} leaves out "
                            f"{col_range(b.name, col, missing).split('!')[1]}")
                continue
            expected = sum(number(b.value(x, col)) for x in with_number)
            actual = b.value(r, col)
            if with_number and isinstance(actual, (int, float)) and not isinstance(actual, bool) \
                    and abs(actual - expected) > max(0.01, abs(expected) * 1e-6):
                deviations.append(f"{b.name}!{cell.coordinate} = {format_number(float(actual))}, "
                                  f"the data rows give {format_number(expected)}")
    for col in total_cols:  # total column: =SUM(B2:G2) per row
        for r in b.data + b.totals:
            cell = ws.cell(r, col)
            if not is_formula(cell):
                continue
            g = _single_range(str(cell.value))
            if not g or g[1] != g[3] or g[1] != r or g[2] >= col:
                continue
            lab = b.label(r)
            start = (lab[0] + 1) if lab else 1
            with_number = [c for c in range(start, col) if c not in total_cols and number(b.value(r, c)) is not None]
            # left of the range only columns of the same kind count: area, share or head count in front of the
            # cost columns are the key of an allocation, not a forgotten summand
            missing = [c for c in with_number if c > g[2] or (c < g[0] and _same_kind(b, r, c, range(g[0], g[2] + 1)))]
            if missing:
                gaps.append(f"{b.name}!{cell.coordinate} ={str(cell.value)[1:40]} leaves out "
                            f"{', '.join(col_letter(c) + str(r) for c in missing)}")
    parts = gaps + deviations + others
    if parts:
        ctx.report("sum_range", "sum incomplete or wrong: " + "; ".join(parts[:3])
                   + (f" and {len(parts) - 3} more" if len(parts) > 3 else ""))


# --------------------------------------------------------------------------------------- task checks --

def _evaluate_formula(formula: str, values: dict) -> float | None:
    """Value of a [checks.calculations] formula ('menge * preis') for the numbers of one row."""
    try:
        tree = ast.parse(formula, mode="eval").body
    except SyntaxError:
        return None

    class Substitute(ast.NodeTransformer):
        def visit_Name(self, node):
            return ast.copy_location(ast.Constant(values[node.id]), node)

    try:
        return evaluate(Substitute().visit(tree))
    except (ZeroDivisionError, ValueError, KeyError):
        return None


def _expected_display(expected: float, like) -> str:
    """Expected value in the form of the cell (a date for date cells)."""
    if isinstance(like, (datetime, date)):
        return (datetime(1899, 12, 30) + timedelta(days=round(expected))).strftime("%d.%m.%Y")
    return format_number(expected)


def _calculation(sheets: list, r: dict, ctx: Context) -> None:
    """One entry of [[checks.calculations]]: find the columns (or rows of a transposed table), redo the formula."""
    name, formula, result = r["name"], r["formula"], r["result"]
    tolerance = float(r.get("tolerance", 0.01))
    names = list(r["columns"])
    inputs = [n for n in names if n != result]
    additive = not re.search(r"[*/]", formula)
    chosen = sheets
    if "sheet" in r:
        if isinstance(r["sheet"], int):
            chosen = sheets[r["sheet"] - 1:r["sheet"]]
        else:
            chosen = [b for b in sheets if re.search(r["sheet"], b.name, re.I)]
    not_found = None  # names not found in the best sheet/orientation ([] = all found somewhere)
    for b in chosen:
        for across in (False, True):
            if across:  # table across: names are row labels, values run along the columns
                cand = {n: [z for z in b.data + b.totals if (lab := b.label(z)) and re.search(m, lab[1], re.I)]
                        for n, m in r["columns"].items()}
                label_cols = {lab[0] for z in b.data + b.totals if (lab := b.label(z))}
                start = max(label_cols) + 1 if label_cols else 2
                vectors = [(c, TOTAL_LABEL.match(b.header(c)) is not None) for c in range(start, b.width + 1)]
            else:
                cand = {n: [c for c, head in b.headers.items() if re.search(m, head, re.I)]
                        for n, m in r["columns"].items()}
                vectors = [(z, False) for z in b.data] + [(z, True) for z in b.totals]
            absent = [n for n in names if not cand[n]]
            if absent:
                if not_found is None or len(absent) < len(not_found):
                    not_found = absent
                continue
            not_found = []
            best = None
            for assignment in itertools.islice(itertools.product(*(cand[n] for n in names)), 200):
                if len(set(assignment)) < len(assignment):
                    continue
                at = dict(zip(names, assignment, strict=True))
                result_head = (b.label(at[result]) or (0, ""))[1] if across else b.header(at[result])
                percent_head = bool(re.search(r"%|prozent", result_head, re.I))
                outcome = {"ok": 0, "wrong": [], "missing": [], "total_missing": [], "total_without": []}
                for v, is_total in vectors:
                    if is_total and not additive:
                        continue
                    # (row, column) of each name in this vector; cells are inside the sheet, so creating them in
                    # openpyxl does not change its size
                    rc = {n: (at[n], v) if across else (v, at[n]) for n in names}
                    coord = {n: f"{col_letter(c)}{r}" for n, (r, c) in rc.items()}
                    cell = {n: b.ws.cell(*rc[n]) for n in names}
                    raw = {n: b.value(*rc[n]) for n in names}
                    nums = {n: number_or_date(w) for n, w in raw.items()}
                    present = [n for n in inputs if nums[n] is not None]
                    if is_total:
                        is_sum = {n: is_formula(cell[n]) and bool(SUM.search(str(cell[n].value))) for n in names}
                        # a total row is only checked where it adds up (SUM); MIN/MAX/AVERAGE totals are no sums
                        if all(is_sum[n] for n in inputs) and is_sum[result]:
                            pass
                        elif all(is_sum[n] for n in inputs) and empty(raw[result]):
                            outcome["total_missing"].append(coord[result])
                            continue
                        elif is_sum[result] and present != inputs and all(empty(raw[n]) for n in inputs
                                                                          if nums[n] is None):
                            outcome["total_without"] += [coord[n] for n in inputs if nums[n] is None]
                            continue
                        else:
                            continue
                    if len(present) == len(inputs) and nums[result] is not None:
                        expected = _evaluate_formula(formula, nums)
                        if expected is None:
                            continue
                        # 'Sparquote in %' = 17,31 without percent format: the share × 100 is right as well
                        scale = percent_head and not percent_format(cell[result].number_format)
                        if abs(expected - nums[result]) <= tolerance or (
                                scale and abs(100 * expected - nums[result]) <= 100 * tolerance):
                            outcome["ok"] += 1
                        else:
                            outcome["wrong"].append(f"{b.name}!{coord[result]} = {display(raw[result])}, "
                                                    f"expected {_expected_display(expected, raw[result])}")
                    elif len(present) == len(inputs) and empty(raw[result]) and not is_total:
                        outcome["missing"].append(coord[result])
                score = (outcome["ok"], -len(outcome["wrong"]))
                if best is None or score > best[0]:
                    best = (score, outcome)
            if best is None:
                continue
            e = best[1]
            if not (e["ok"] or e["wrong"] or e["missing"]):
                continue
            if e["wrong"]:
                ctx.report("calculations", f"calculation „{name}“ is wrong: {'; '.join(e['wrong'][:3])}"
                                           + (f" and {len(e['wrong']) - 3} more" if len(e["wrong"]) > 3 else ""))
            if e["missing"]:
                ctx.report("calculations", f"calculation „{name}“: {r['result']} missing in "
                                           f"{b.name}!{', '.join(e['missing'][:6])}", label="missing result")
            if e["total_missing"] or e["total_without"]:
                cells = e["total_missing"] + e["total_without"]
                ctx.report("calculations", f"total row incomplete (calculation „{name}“): no value in "
                                           f"{b.name}!{', '.join(cells[:6])}", label="incomplete total row",
                           severity=SOFT)
            return
    if not_found:
        ctx.not_checkable("calculations", f"calculation „{name}“: no column/row for "
                                          + ", ".join(f"{n} („{r['columns'][n]}“)" for n in not_found))
    else:
        ctx.not_checkable("calculations", f"calculation „{name}“: no row with all values")


def _column_values(sheets: list, pattern: str):
    """[(sheet, header, [(coordinate, value)])] of all columns (or row-labelled rows) matching the pattern."""
    out = []
    for b in sheets:
        for c, head in b.headers.items():
            if re.search(pattern, head, re.I):
                out.append((b, head, [(f"{col_letter(c)}{r}", b.value(r, c)) for r in b.data]))
        if not out:
            for r in b.data:
                lab = b.label(r)
                if lab and re.search(pattern, lab[1], re.I):
                    out.append((b, lab[1], [(f"{col_letter(c)}{r}", b.value(r, c))
                                            for c in range(lab[0] + 1, b.width + 1)]))
    return out


# Columns whose dates must lie in dates_between (hard); other date cells are only soft, dates in text are ignored.
DATE_HEADER = re.compile(r"datum|\bdate\b|tag\b|kauf|anschaff|gekauft|start|beginn|\bende\b|termin|frist|fällig|"
                         r"abfahrt|reise", re.I)
VALIDITY_HEADER = re.compile(r"\bbis\b|garantie|ablauf|gültig|nächste|wartung|ende der", re.I)


def _is_date_column(b: Sheet, c: int) -> bool:
    """A column with dates: the header names a date (Datum, Start, Bereitschaftstag …) or most filled data cells hold
    a date (date number format, also as a formula result)."""
    if DATE_HEADER.search(b.headers.get(c, "")):
        return True
    values = [v for r in b.data if not empty(v := b.value(r, c))]
    return len(values) >= 2 and sum(1 for v in values if isinstance(v, (datetime, date))) >= len(values) / 2


def _data_rows(sheets, ctx):
    first, p = sheets[0] if sheets else None, ctx.checks
    n = sum(1 for r in first.data if sum(1 for c in range(1, first.width + 1)
                                         if not empty(first.ws.cell(r, c).value)) >= 2) if first else 0
    if n < p["data_rows_min"]:
        ctx.report("data_rows_min", f"only {n} data rows on the first sheet (expected ≥ {p['data_rows_min']})")


def _months(sheets, ctx):
    lo, hi = ctx.checks["months"]
    months = set()
    for b in sheets:
        for row in b.ws.iter_rows():
            for c in row:
                months |= months_in(c.value, lo, hi)
    if len(months) < hi - lo + 1:
        found = ", ".join(SHORT_MONTH[n] for n in sorted(months)) or "none"
        ctx.report("months", f"months {MONTH_NAMES[lo]}–{MONTH_NAMES[hi]} incomplete: {found} "
                             f"({len(months)} of {hi - lo + 1})")


def _month_rows(sheets, ctx):
    expected = ctx.checks["month_rows_min"]
    max_rows = max((len(b.data) for b in sheets), default=0)
    month_cols = max((sum(1 for head in b.headers.values() if months_in(head)) for b in sheets), default=0)
    if max_rows < expected and month_cols < expected:
        ctx.report("month_rows_min", f"fewer than {expected} month rows or columns (at most {max_rows} data rows)")


def _dates_between(sheets, ctx):
    """Dates in date columns (header Datum, Kauf, Start, Ende, Termin, Stichtag …) outside the range: hard; date
    cells in other columns ('Erledigt am' may lie after the period): soft (not in validity columns like 'Garantie
    bis'); dates in running text are not counted."""
    lo, hi = (date.fromisoformat(x) for x in ctx.checks["dates_between"])
    found, hard, soft = 0, [], []
    for b in sheets:
        for c in range(1, b.width + 1):
            head = b.headers.get(c, "")
            date_col = bool(DATE_HEADER.search(head)) and not VALIDITY_HEADER.search(head)
            if VALIDITY_HEADER.search(head):
                continue
            for r in b.data:
                v = b.value(r, c)
                if isinstance(v, datetime):
                    dates = [v.date()]
                elif isinstance(v, date):
                    dates = [v]
                elif isinstance(v, str) and date_col:
                    dates = [d for f in dates_in(v) if f.year and (d := f.with_year(f.year))]
                else:
                    dates = []
                for d in dates:
                    found += 1
                    if not lo <= d <= hi:
                        (hard if date_col else soft).append(f"{b.name}!{col_letter(c)}{r} = {d:%d.%m.%Y}")
    range_text = f"{lo:%d.%m.}–{hi:%d.%m.%Y}"
    if not found:
        ctx.not_checkable("dates_between", f"no dates in the data rows (expected {range_text})")
    for items, severity in ((hard, HARD), (soft, SOFT)):
        if items:
            where = "" if severity == HARD else " (no date column)"
            ctx.report("dates_between", f"date outside {range_text}{where}: {', '.join(items[:3])}"
                                        + (f" and {len(items) - 3} more" if len(items) > 3 else ""), severity=severity)


def _columns(sheets, ctx):
    texts = [head for b in sheets for head in b.headers.values()]
    texts += [lab[1] for b in sheets for r in b.data + b.totals if (lab := b.label(r))]
    missing = [m for m in ctx.checks["columns"] if not any(re.search(m, t, re.I) for t in texts)]
    if missing:
        ctx.report("columns", "required columns missing: " + ", ".join(f"„{m}“" for m in missing))


def _total_row(sheets, ctx):
    """A SUM in a total row of the first sheet (or in a total column of a table across); running sums do not count."""
    first = sheets[0] if sheets else None
    if first is None:
        return
    total_cols = {c for c, head in first.headers.items() if TOTAL_LABEL.match(head)}
    has = any(SUM.search(str(c.value)) and (c.row in first.totals or c.column in total_cols)
              for c in first.formulas())
    if not has:
        ctx.report("total_row", "no total row (SUM formula) on the first sheet")


def _line_item_formula(sheets, ctx):
    first = sheets[0] if sheets else None
    if first and not any(re.search(r"\*|PRODUCT|PRODUKT", str(c.value), re.I) for c in first.formulas()
                         if c.row in first.data):
        ctx.report("line_item_formula", "no formula for the line total (quantity × unit price)")


def _priced_items(sheets, ctx):
    b = sheets[0] if sheets else None
    if b is None:
        return
    price = [col for col, head in b.headers.items()
             if re.search(r"preis|betrag|€|\beur\b|kosten", head, re.I)
             and not re.search(r"gesamt|summe|total|netto|brutto", head, re.I)]
    if not price:
        price = [col for col in b.headers if any(is_number_format(b.ws.cell(r, col).number_format) for r in b.data)
                 and not any(is_formula(b.ws.cell(r, col)) for r in b.data)]
    with_price = sum(1 for r in b.data if any(isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0
                                              for v in (b.value(r, col) for col in price)))
    expected = ctx.checks["priced_items_min"]
    if with_price < expected:
        ctx.report("priced_items_min", f"only {with_price} line items with a unit price > 0 (expected {expected})")


def _value_range(sheets, w, ctx):
    columns = _column_values(sheets, w["column"])
    if not columns:
        ctx.not_checkable("value_ranges", f"no column „{w['column']}“ for the value range")
        return
    lo, hi = w.get("min", float("-inf")), w.get("max", float("inf"))
    outside = [f"{b.name}!{coord} = {format_number(x)}" for b, _, values in columns for coord, v in values
               if (x := number(v)) is not None and not lo <= x <= hi]
    if outside:
        bounds = (f"{format_number(lo) if lo != float('-inf') else '…'}–"
                  f"{format_number(hi) if hi != float('inf') else '…'}")
        ctx.report("value_ranges", f"value outside {bounds} in „{columns[0][1]}“: {', '.join(outside[:3])}"
                                   + (f" and {len(outside) - 3} more" if len(outside) > 3 else ""))


def _sorted(sheets, ctx):
    """The numbers (or dates) of the column in the data rows of the first sheet that has it are sorted; with several
    matching columns one sorted column is enough. `sheet` (position from 1 or regex on the name) limits the
    sheets."""
    s = ctx.checks["sorted"]
    descending = bool(s.get("descending"))
    direction = "descending" if descending else "ascending"
    if "sheet" in s:
        sheets = (sheets[s["sheet"] - 1:s["sheet"]] if isinstance(s["sheet"], int)
                  else [b for b in sheets if re.search(s["sheet"], b.name, re.I)])
    for b in sheets:
        cols = [c for c, head in b.headers.items() if re.search(s["column"], head, re.I)]
        if not cols:
            continue
        breaks = []
        for c in cols:
            values = [(r, x) for r in b.data if (x := number_or_date(b.value(r, c))) is not None]
            if len(values) < 2:
                continue
            brk = next(((a, z) for a, z in zip(values, values[1:], strict=False)
                        if (z[1] > a[1] if descending else z[1] < a[1])), None)
            if brk is None:
                return
            breaks.append((c, brk))
        if not breaks:
            ctx.not_checkable("sorted", f"no numbers in column „{b.header(cols[0])}“ ({b.name})")
            return
        c, ((r1, _), (r2, _)) = breaks[0]
        ctx.report("sorted", f"„{b.header(c)}“ not sorted {direction}: {b.name}!{col_letter(c)}{r2} = "
                             f"{display(b.value(r2, c))} after {col_letter(c)}{r1} = {display(b.value(r1, c))}")
        return
    where = f" on sheet „{s['sheet']}“" if "sheet" in s else ""
    ctx.not_checkable("sorted", f"no column „{s['column']}“{where} to check the order")


def _sheet(sheets: list, sheet):
    """Sheet by position (from 1) or by a regex on its name."""
    if isinstance(sheet, int):
        return sheets[sheet - 1] if 0 < sheet <= len(sheets) else None
    return next((b for b in sheets if re.search(sheet, b.name, re.I)), None)


def _column_with(b: Sheet, pattern: str, exclude=(), text: bool = False) -> int | None:
    """Column whose header matches; of several the one with the most text (text=True) or number values."""
    cand = [c for c, head in b.headers.items() if c not in exclude and re.search(pattern, head, re.I)]

    def hits(c):
        values = [b.value(r, c) for r in b.data]
        if text:
            return sum(1 for v in values if isinstance(v, str) and v.strip() and number(v) is None)
        return sum(1 for v in values if number(v) is not None)
    return max(cand, key=lambda c: (hits(c), -c)) if cand else None


def _group_norm(v) -> str:
    return re.sub(r"[\W_]+", "", (v if isinstance(v, str) else display(v)).lower())


def _to_group(n: str, groups: dict) -> str | None:
    """Source group of an overview row: equal, else one name a prefix of the other with ≥ 4 letters (Laptop/Laptops)."""
    if n in groups:
        return n
    cand = [g for g in groups if min(len(g), len(n)) >= 4 and (g.startswith(n) or n.startswith(g))]
    return min(cand, key=lambda g: abs(len(g) - len(n))) if cand else None


def _groups(sheets, g, ctx):
    """One [[checks.groups]]: count and sum per group in the source sheet, compared with the overview sheet."""
    q, z = g["source"], g["target"]
    bq, bz = _sheet(sheets, q["sheet"]), _sheet(sheets, z["sheet"])
    if bq is None or bz is None:
        ctx.not_checkable("groups", f"summary „{g['name']}“: no sheet {q['sheet'] if bq is None else z['sheet']}")
        return
    gq = _column_with(bq, q["group"], text=True)
    vq = _column_with(bq, q["value"], {gq}) if "value" in q else None
    gz = _column_with(bz, z["group"], text=True)
    cz = _column_with(bz, z["count"], {gz}) if "count" in z else None
    vz = _column_with(bz, z["value"], {gz, cz}) if "value" in z else None
    missing = [f"{what} „{m}“ ({b.name})" for what, m, b, c in (
        ("group", q["group"], bq, gq), ("value", q.get("value"), bq, vq), ("group", z["group"], bz, gz),
        ("count", z.get("count"), bz, cz), ("value", z.get("value"), bz, vz)) if m is not None and c is None]
    if missing:
        ctx.not_checkable("groups", f"summary „{g['name']}“: no column {', '.join(missing)}")
        return
    groups = {}  # normalised name -> [name, count, sum]
    for r in bq.data:
        v = bq.value(r, gq)
        if empty(v) or not (n := _group_norm(v)):
            continue
        entry = groups.setdefault(n, [display(v).strip(), 0, 0.0])
        entry[1] += 1
        entry[2] += (number(bq.value(r, vq)) or 0.0) if vq else 0.0
    if not groups:
        ctx.not_checkable("groups", f"summary „{g['name']}“: no groups in column „{bq.header(gq)}“ ({bq.name})")
        return
    deviations, seen = [], set()
    for r in bz.data:
        v = bz.value(r, gz)
        if empty(v) or (n := _to_group(_group_norm(v), groups)) is None:
            continue
        seen.add(n)
        name, count, total = groups[n]
        for c, expected, tolerance, what in ((cz, count, 0.0, "count"),
                                             (vz, total, max(0.01, abs(total) * 0.001), "value")):
            if c is None:
                continue
            actual = number(bz.value(r, c))
            if actual is None or abs(actual - expected) > tolerance + 1e-9:
                shown = "empty" if actual is None else format_number(actual)
                deviations.append(f"{display(v).strip()}: {what} {shown} ({bz.name}!{col_letter(c)}{r}), "
                                  f"per {bq.name} {format_number(expected)}")
    deviations += [f"„{name}“ missing" for n, (name, _, _) in groups.items() if n not in seen]
    if deviations:
        ctx.report("groups", f"summary „{g['name']}“ does not match {bq.name}: {'; '.join(deviations[:3])}"
                             + (f" and {len(deviations) - 3} more" if len(deviations) > 3 else ""))


def task_checks(sheets: list, charts: list, n_formulas: int, ctx: Context) -> None:
    """The xlsx keys of [checks]; each one guarded (an exception only gives a soft defect)."""
    p = ctx.checks
    if "sheets_min" in p and len(sheets) < p["sheets_min"]:
        text = "no second sheet" if p["sheets_min"] == 2 else f"only {len(sheets)} sheets"
        ctx.report("sheets_min", f"{text} (expected ≥ {p['sheets_min']} sheets)")
    if "formulas_min" in p and n_formulas < p["formulas_min"]:
        ctx.report("formulas_min", "no formula" if n_formulas == 0
                   else f"only {n_formulas} formulas (expected ≥ {p['formulas_min']})")
    steps = [("data_rows_min", _data_rows), ("months", _months), ("month_rows_min", _month_rows),
             ("dates_between", _dates_between), ("columns", _columns), ("total_row", _total_row),
             ("line_item_formula", _line_item_formula), ("priced_items_min", _priced_items), ("sorted", _sorted)]
    for key, step in steps:
        if p.get(key):
            with ctx.guarded(key):
                step(sheets, ctx)
    with ctx.guarded("chart"):
        check_charts(charts, ctx)
    for r in p.get("calculations", []):
        with ctx.guarded("calculations"):
            _calculation(sheets, r, ctx)
    for w in p.get("value_ranges", []):
        with ctx.guarded("value_ranges"):
            _value_range(sheets, w, ctx)
    for g in p.get("groups", []):
        with ctx.guarded("groups"):
            _groups(sheets, g, ctx)
    if "column_format" in p:
        with ctx.guarded("column_format"):
            _column_format(sheets, ctx)
    if p.get("workdays"):
        with ctx.guarded("workdays"):
            _workdays(sheets, ctx)


FORMATS = {"percent": "percent format", "date": "date format", "euro": "euro format"}


def _has_format(kind: str, cell) -> bool:
    from openpyxl.styles.numbers import is_date_format

    fmt = cell.number_format or ""
    if kind == "percent":
        return percent_format(fmt)
    if kind == "date":
        return isinstance(cell.value, (datetime, date)) or is_date_format(fmt)
    return "€" in fmt or "EUR" in fmt.upper()


def _format_cells(sheets, pattern: str):
    """[(sheet, header, [(row, column)])] of the data cells under matching headers; in a sheet without such a column
    the cells right of a matching row label (table across)."""
    out = []
    for b in sheets:
        cols = [(c, head) for c, head in b.headers.items() if re.search(pattern, head, re.I)]
        out += [(b, head, [(r, c) for r in b.data]) for c, head in cols]
        if not cols:
            out += [(b, lab[1], [(r, c) for c in range(lab[0] + 1, b.width + 1)])
                    for r in b.data if (lab := b.label(r)) and re.search(pattern, lab[1], re.I)]
    return out


def _column_format(sheets, ctx):
    """column_format = { regex on the header = "percent" | "date" | "euro" }: the values of the data rows in these
    columns (or rows of a table across) have the number format; numbers and dates typed as text do not."""
    for pattern, kind in ctx.checks["column_format"].items():
        areas = _format_cells(sheets, pattern)
        if not areas:
            ctx.not_checkable("column_format", f"no column „{pattern}“ for the {FORMATS[kind]}")
            continue
        wrong, seen = [], 0
        for b, head, cells in areas:
            for r, c in cells:
                cell, value = b.ws.cell(r, c), b.value(r, c)
                if empty(value):
                    continue
                text = isinstance(value, str)
                if text and not (number(value) is not None or dates_in(value)):
                    continue  # a note like 'laufend' is no value
                seen += 1
                if text or not _has_format(kind, cell):
                    wrong.append((b, head, c, r, value, text))
        if wrong:
            b, head, c, r, value, text = wrong[0]
            example = f"„{value}“ as text" if text else f"{display(value)} (format „{b.ws.cell(r, c).number_format}“)"
            ctx.report("column_format", f"„{head}“ not in {FORMATS[kind]}: {len(wrong)} of {seen} values, e.g. "
                                        f"{b.name}!{col_letter(c)}{r} = {example}")


def easter(year: int) -> date:
    """Easter Sunday (Gregorian calendar, anonymous algorithm)."""
    a, b, c = year % 19, year // 100, year % 100
    d, e = divmod(b, 4)
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, kk = divmod(c, 4)
    l_ = (32 + 2 * e + 2 * i - h - kk) % 7
    m = (a + 11 * h + 22 * l_) // 451
    month, day = divmod(h + l_ - 7 * m + 114, 31)
    return date(year, month, day + 1)


def holidays_bw(year: int) -> dict:
    """Public holidays in Baden-Württemberg: {date: German name}."""
    o = easter(year)
    return {date(year, 1, 1): "Neujahr", date(year, 1, 6): "Heilige Drei Könige", o - timedelta(days=2): "Karfreitag",
            o + timedelta(days=1): "Ostermontag", date(year, 5, 1): "Tag der Arbeit",
            o + timedelta(days=39): "Christi Himmelfahrt", o + timedelta(days=50): "Pfingstmontag",
            o + timedelta(days=60): "Fronleichnam", date(year, 10, 3): "Tag der Deutschen Einheit",
            date(year, 11, 1): "Allerheiligen", date(year, 12, 25): "1. Weihnachtstag",
            date(year, 12, 26): "2. Weihnachtstag"}


# Columns with dates that are no working days by nature, and rows about days off.
NO_WORKDAY_HEADER = re.compile(r"geburt|kauf|anschaff|garantie|gültig|ablauf|feiertag|ferien|erstellt|stand", re.I)
DAY_OFF = re.compile(r"feiertag|wochenende|samstag|sonntag|ferien|betriebsurlaub|geschlossen|\bfrei\b", re.I)
# A cell that only names the weekday (column „Wochentag“): it does not make the row a day off.
WEEKDAY_CELL = re.compile(r"(mo|di|mi|do|fr|sa|so)\.?|(montag|dienstag|mittwoch|donnerstag|freitag|samstag|sonntag)",
                          re.I)
WEEKDAY_NAMES = ["Montag", "Dienstag", "Mittwoch", "Donnerstag", "Freitag", "Samstag", "Sonntag"]


def _as_date(v, head: str) -> date | None:
    """A date cell, or a serial number (2000–2099) in a column whose header names a date."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool) and DATE_HEADER.search(head) and 36526 <= v < 73051:
        return date(1899, 12, 30) + timedelta(days=int(v))
    return None


def _workdays(sheets, ctx):
    """workdays = true (Saturday/Sunday) or "BW" (also public holidays in Baden-Württemberg): dates in the data rows
    lie on working days. Day-by-day calendars (three days in a row across a weekend in one column), columns like
    'Kaufdatum' or 'Feiertag' and rows about days off (a weekday name alone in its cell does not count) are left
    out."""
    with_holidays = ctx.checks["workdays"] == "BW"
    found, seen = [], 0
    for b in sheets:
        if DAY_OFF.search(b.name):
            continue
        for c in range(1, b.width + 1):
            head = b.headers.get(c, "")
            if NO_WORKDAY_HEADER.search(head):
                continue
            dates = [(r, d) for r in b.data if (d := _as_date(b.value(r, c), head))]
            days = sorted({d for _, d in dates})
            if any(days[i + 2] - days[i] == timedelta(days=2) and any(t.weekday() >= 5 for t in days[i:i + 3])
                   for i in range(len(days) - 2)):
                continue  # a calendar with one row per day, weekends included (Mon–Fri lists are checked)
            for r, d in dates:
                row = " ".join(str(v) for x in range(1, b.width + 1)
                               if (v := b.value(r, x)) is not None and not WEEKDAY_CELL.fullmatch(str(v).strip()))
                if DAY_OFF.search(row):
                    continue
                seen += 1
                off = WEEKDAY_NAMES[d.weekday()] if d.weekday() >= 5 else (
                    holidays_bw(d.year).get(d) if with_holidays else None)
                if off:
                    found.append(f"{b.name}!{col_letter(c)}{r} = {d:%d.%m.%Y} ({off})")
    if found:
        kind = "a weekend or public holiday" if with_holidays else "a weekend"
        ctx.report("workdays", f"{len(found)} of {seen} dates on {kind}: {', '.join(found[:3])}"
                               + (f" and {len(found) - 3} more" if len(found) > 3 else ""))


# ----------------------------------------------------------------------------------------------- run --

def check_xlsx(path: Path, ctx: Context, recalculated: dict) -> tuple[dict, Document]:
    """Checks one workbook; returns metrics and the text document for the text checks. `recalculated` maps the
    original path to the LibreOffice copy (from recalculate()); without one, formula values are unknown."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    values_path = recalculated.get(path)
    wbv = load_workbook(values_path, data_only=True) if values_path else None
    if wbv is None:
        ctx.report("not_recalculated", "LibreOffice could not open/recalculate the file")
    sheets = [Sheet(ws, wbv[ws.title] if wbv is not None and ws.title in wbv.sheetnames else None)
              for ws in wb.worksheets]
    doc = Document(numbers=[], constants=[])
    n_formulas = 0
    for b in sheets:
        n_formulas += sum(1 for _ in b.formulas())
        general_checks(b, ctx)
        rows = []
        for row in b.ws.iter_rows():
            shown = []
            for c in row:
                v = c.value
                if isinstance(v, str) and not is_formula(c) and v.strip():
                    doc.paragraphs.append((f"{b.name}!{c.coordinate}", v))
                elif isinstance(v, (datetime, date)):
                    doc.paragraphs.append((f"{b.name}!{c.coordinate}", display(v)))
                value = b.value(c.row, c.column) if is_formula(c) else v
                if isinstance(v, (int, float)) and not isinstance(v, bool) and not is_formula(c):
                    doc.constants.append((f"{b.name}!{c.coordinate}", float(v), percent_format(c.number_format)))
                if (x := number(value)) is not None:  # values_present: numbers (also as text); 12,5 % also as 12.5
                    percent = percent_format(c.number_format) and not isinstance(value, str)
                    doc.numbers += [x, 100 * x] if percent else [x]
                shown.append(display(value))
            if any(shown):
                rows.append(shown)
        doc.tables.append(Table(b.name, rows, calc=False))
    charts, resolver = [], _resolver(sheets)
    with zipfile.ZipFile(path) as z:
        for n in sorted(z.namelist()):
            if re.match(r"xl/charts/chart\d+\.xml$", n):
                with ctx.guarded("chart"):
                    charts.append(chart_info(z.read(n), resolver))
    task_checks(sheets, charts, n_formulas, ctx)
    metrics = {"sheets": len(sheets), "data_rows": sum(len(b.data) for b in sheets), "formulas": n_formulas,
               "charts": len(charts)}
    return metrics, doc
