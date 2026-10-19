"""What the quality judge sees of a built file: Markdown text and page images.

    render_text(path, recalculated=None) -> str     Markdown of a docx/xlsx/pptx (no LibreOffice)
    render_texts(paths) -> {path: str}              many files; all xlsx recalculated in one LibreOffice run
    render_images(paths, out_dir) -> {path: [png]}  page images via LibreOffice -> PDF -> pdftoppm
    combine_pages(pngs, dst, columns)               the pages of one file as one image
    neutral_copy(path, dst)                         copy without build times (ZIP entries, core.xml, embedded workbooks)

The text gives the judge what an image cannot show reliably: formulas next to their results, numbers stored as text,
number formats, list numbering as Word shows it, chart series with their cell ranges, speaker notes. The images
show what the text cannot: layout, overflowing text, charts as drawn. Structural labels are English (the judge prompt
is English), the content stays German as written.

What the text contains:
- docx: header/footer (page fields as [Page]/[Page count]), title, headings, list numbering as Word shows it (1., a),
  bullets, levels), bold runs, tables (merged cells once), TOC field, page breaks, pictures, text boxes, charts.
- xlsx: every sheet as a table with column letters and row numbers, formulas with their recalculated result, text
  that looks like a number or formula marked as text, number formats, merged cells, print header/footer, comments,
  charts (type, title, anchor cell, series with ranges and resolved values).
- pptx: per slide layout, title/subtitle, text in reading order with levels, tables, charts (type, categories,
  series), pictures, shapes reaching over the slide edge, speaker notes.
"""

import io
import posixpath
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date, datetime, time
from pathlib import Path

from lxml import etree
from openpyxl.utils import get_column_letter

from officebench.checks.xlsx import SOFFICE, _profile, _run_soffice, recalculate

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "v": "urn:schemas-microsoft-com:vml",
}
_PARSER = etree.XMLParser(resolve_entities=False, no_network=True)
FILE_TYPES = (".docx", ".xlsx", ".pptx")

# Images: docx/pptx pages at DPI (an A4 page is about 370 x 530 px, a 16:9 slide 600 x 340 px: enough to judge layout,
# the text comes from render_text); every xlsx sheet on one page, scaled so the longer side has XLSX_SIZE px.
DPI = 45
XLSX_SIZE = 1100
MAX_PAGES = {".docx": 8, ".pptx": 12, ".xlsx": 12}  # Word pages, slides, sheets
COLUMNS = {".docx": 4, ".pptx": 3, ".xlsx": 1}  # combine_pages: pages per row
# Calc's PDF export with one page per sheet, so no chart is cut in half by a page break.
XLSX_PDF = 'pdf:calc_pdf_Export:{"SinglePageSheets":{"type":"boolean","value":"true"}}'


def _q(tag: str) -> str:
    prefix, name = tag.split(":")
    return f"{{{NS[prefix]}}}{name}"


def _local(el) -> str:
    return etree.QName(el).localname if isinstance(el.tag, str) else ""


# ---------------------------------------------------------------------------------------------------- common helpers

def _cell(text) -> str:
    return str(text).replace("\r", "").replace("\n", " / ").replace("|", "\\|").strip()


def md_table(rows: list[list]) -> str:
    """Markdown table; the first row is the header row, short rows are padded."""
    rows = [r for r in rows if r is not None]
    width = max((len(r) for r in rows), default=0)
    if not width:
        return ""
    rows = [[_cell(c) for c in r] + [""] * (width - len(r)) for r in rows]
    return "\n".join(["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width] +
                     ["| " + " | ".join(r) + " |" for r in rows[1:]])


def _value(v) -> str:
    """A cell value as text; dates in German order, since that is what the German files show."""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() and abs(v) < 1e15 else f"{v:.10g}"
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y %H:%M") if (v.hour, v.minute, v.second) != (0, 0, 0) else v.strftime("%d.%m.%Y")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, time):
        return v.strftime("%H:%M")
    return str(v)


# --------------------------------------------------------------------------------------------------------- charts

CHART_TYPES = {
    "barChart": "column chart", "bar3DChart": "3D column chart", "lineChart": "line chart",
    "line3DChart": "3D line chart", "pieChart": "pie chart", "pie3DChart": "3D pie chart",
    "doughnutChart": "doughnut chart", "ofPieChart": "pie-of-pie chart", "areaChart": "area chart",
    "area3DChart": "3D area chart", "scatterChart": "scatter chart (XY)", "radarChart": "radar chart",
    "bubbleChart": "bubble chart", "stockChart": "stock chart", "surfaceChart": "surface chart",
    "surface3DChart": "3D surface chart",
}
AXES = (("c:catAx", "category axis"), ("c:dateAx", "date axis"), ("c:valAx", "value axis"), ("c:serAx", "series axis"))


def _texts(el) -> str:
    return "".join(t.text or "" for t in el.iter(_q("a:t"))) if el is not None else ""


def _cache(data) -> list:
    """Values of a c:strCache / c:numCache / literal in point order (missing points = None)."""
    if data is None:
        return []
    count = data.find(_q("c:ptCount"))
    points = {int(p.get("idx", 0)): (p.findtext(_q("c:v")) or "") for p in data.findall(_q("c:pt"))}
    if count is not None and count.get("val", "").isdigit():
        n = int(count.get("val"))
    else:
        n = max(points) + 1 if points else 0
    return [points.get(i) for i in range(n)]


def _data(el, resolve) -> tuple[str | None, list]:
    """(range formula, values) of c:tx / c:cat / c:val / c:xVal / c:yVal; the resolved cells if possible, since
    openpyxl writes no cache and a stale cache would show values the sheet does not have."""
    if el is None:
        return None, []
    for ref_tag, cache_tag in (("c:numRef", "c:numCache"), ("c:strRef", "c:strCache"),
                               ("c:multiLvlStrRef", "c:multiLvlStrCache")):
        ref = el.find(_q(ref_tag))
        if ref is not None:
            formula = ref.findtext(_q("c:f"))
            values = resolve(formula) if (resolve and formula) else None
            if values is None:
                cache = ref.find(_q(cache_tag))
                if ref_tag == "c:multiLvlStrRef" and cache is not None:
                    cache = cache.find(_q("c:lvl"))
                values = _cache(cache)
            return formula, values
    for lit in ("c:numLit", "c:strLit"):
        if (x := el.find(_q(lit))) is not None:
            return None, _cache(x)
    if (v := el.find(_q("c:v"))) is not None:
        return None, [v.text or ""]
    return None, []


def _point(v) -> str:
    if v is None or v == "":
        return "–"
    if isinstance(v, str):
        return _value(float(v)) if re.fullmatch(r"-?\d+(\.\d+)?([eE][-+]?\d+)?", v) else v
    return _value(v)


def _points(values: list, limit: int = 40) -> str:
    rest = f"; … (+{len(values) - limit})" if len(values) > limit else ""
    return "; ".join(_point(v) for v in values[:limit]) + rest


def _ref(formula: str | None) -> str:
    return f" [{formula}]" if formula else ""


def chart_text(space, resolve=None, where: str = "") -> str:
    """Readable summary of a c:chartSpace element: type, title, categories, series, axis titles, legend.

    resolve(formula) -> list of values or None resolves cell ranges (xlsx, after recalculation); otherwise the values
    cached in the chart part are used (pptx, docx).
    """
    chart = space.find(_q("c:chart")) if _local(space) == "chartSpace" else space
    if chart is None:
        return "**Chart** (not readable)"
    title_el = chart.find(_q("c:title"))
    title = _texts(title_el) or ("(automatic from series name)" if title_el is not None else "")
    plot = chart.find(_q("c:plotArea"))
    kinds, series, categories = [], [], []
    for kind_el in plot if plot is not None else []:
        tag = _local(kind_el)
        if not tag.endswith("Chart"):
            continue
        name = CHART_TYPES.get(tag, tag)
        direction = kind_el.find(_q("c:barDir"))
        if direction is not None and direction.get("val") == "bar":
            name = name.replace("column", "bar")
        grouping = kind_el.find(_q("c:grouping"))
        g = grouping.get("val") if grouping is not None else None
        name += {"stacked": ", stacked", "percentStacked": ", 100% stacked"}.get(g, "")
        details = ", ".join(x for x in (tag, direction.get("val") if direction is not None else None) if x)
        kinds.append(f"{name} ({details})")
        for ser in kind_el.findall(_q("c:ser")):
            _, names = _data(ser.find(_q("c:tx")), resolve)
            cat = ser.find(_q("c:cat")) if ser.find(_q("c:cat")) is not None else ser.find(_q("c:xVal"))
            cat_ref, cats = _data(cat, resolve)
            val = ser.find(_q("c:val")) if ser.find(_q("c:val")) is not None else ser.find(_q("c:yVal"))
            val_ref, values = _data(val, resolve)
            categories.append((cat_ref, cats))
            series.append((_point(names[0]) if names else "(no name)", val_ref, values, cat_ref, cats))
    axes = [f'{label} "{t}"' for tag, label in AXES for ax in (plot.findall(_q(tag)) if plot is not None else [])
            if (t := _texts(ax.find(_q("c:title"))))]
    out = [f"**Chart**{where}: {' + '.join(kinds) or 'no data series'}" + (f'; title "{title}"' if title else
                                                                            "; no title")]
    shared = len({(r, tuple(map(str, k))) for r, k in categories}) == 1
    if shared and categories[0][1]:
        out.append(f"- Categories{_ref(categories[0][0])}: {_points(categories[0][1])}")
    for name, val_ref, values, cat_ref, cats in series:
        line = f'- Series "{name}"{_ref(val_ref)}: {_points(values) if values else "(no values)"}'
        if not shared and cats:
            line += f" | Categories{_ref(cat_ref)}: {_points(cats)}"
        out.append(line)
    if axes:
        out.append("- Axis titles: " + ", ".join(axes))
    if series and chart.find(_q("c:legend")) is None:
        out.append("- no legend")
    return "\n".join(out)


# ----------------------------------------------------------------------------------------------------------- docx

LETTERS = {"lowerLetter": "abcdefghijklmnopqrstuvwxyz", "upperLetter": "ABCDEFGHIJKLMNOPQRSTUVWXYZ"}
ROMAN = [(1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"), (50, "L"), (40, "XL"), (10, "X"),
         (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]


def _roman(n: int) -> str:
    out = ""
    for value, sign in ROMAN:
        while n >= value:
            out, n = out + sign, n - value
    return out


def _format_number(n: int, fmt: str) -> str:
    if fmt in LETTERS:
        return LETTERS[fmt][(n - 1) % 26] * ((n - 1) // 26 + 1)
    if fmt in ("lowerRoman", "upperRoman"):
        return _roman(n).lower() if fmt == "lowerRoman" else _roman(n)
    return f"{n:02d}" if fmt == "decimalZero" else str(n)


class _Numbering:
    """List numbers as Word shows them: counters per list, deeper levels restart, startOverride restarts a list. The
    judge has to see "1. 2. 3." where Word shows it, not the bare paragraphs, to judge numbered steps."""

    def __init__(self, doc):
        self.nums: dict[str, tuple[str, dict]] = {}
        self.abstract: dict[str, dict[int, tuple[str, str, int]]] = {}
        self.counters: dict = {}
        try:
            el = doc.part.numbering_part.element
        except Exception:  # noqa: BLE001 - document without numbering part
            return
        for ab in el.findall(_q("w:abstractNum")):
            levels = {}
            for lvl in ab.findall(_q("w:lvl")):
                fmt, text, start = lvl.find(_q("w:numFmt")), lvl.find(_q("w:lvlText")), lvl.find(_q("w:start"))
                levels[int(lvl.get(_q("w:ilvl"), 0))] = (
                    fmt.get(_q("w:val")) if fmt is not None else "decimal",
                    text.get(_q("w:val")) if text is not None else None,
                    int(start.get(_q("w:val"))) if start is not None else 1)
            self.abstract[ab.get(_q("w:abstractNumId"))] = levels
        for num in el.findall(_q("w:num")):
            ab = num.find(_q("w:abstractNumId"))
            starts = [(ov, ov.find(_q("w:startOverride"))) for ov in num.findall(_q("w:lvlOverride"))]
            overrides = {int(ov.get(_q("w:ilvl"), 0)): int(so.get(_q("w:val"))) for ov, so in starts if so is not None}
            self.nums[num.get(_q("w:numId"))] = (ab.get(_q("w:val")) if ab is not None else None, overrides)

    def prefix(self, num_id: str, ilvl: int) -> str | None:
        """'1.', 'a)', '-' (bullet) or None if the list is unknown."""
        if num_id not in self.nums:
            return None
        ab_id, overrides = self.nums[num_id]
        levels = self.abstract.get(ab_id, {})
        if ilvl not in levels:
            return "-" if not levels else None
        fmt, text, _ = levels[ilvl]
        if fmt == "bullet":
            return "-"
        if fmt == "none":
            return ""
        z = self.counters.setdefault(num_id if overrides else ("abstract", ab_id), {})
        z[ilvl] = z.get(ilvl, overrides.get(ilvl, levels[ilvl][2]) - 1) + 1
        for deeper in [k for k in z if k > ilvl]:
            del z[deeper]
        if text is None:
            return f"{z[ilvl]}."

        def replace(m):
            k = int(m.group(1)) - 1
            fmt_k, _, start_k = levels.get(k, ("decimal", None, 1))
            return _format_number(z.get(k, start_k), fmt_k)

        return re.sub(r"%(\d)", replace, text)


class _Field:
    __slots__ = ("instr", "result", "in_result")

    def __init__(self):
        self.instr, self.result, self.in_result = "", [], False


def _field_text(instr: str, result: str) -> str:
    """Page fields as placeholders (their result is whatever page Word last computed); a TOC with its shown text,
    so the judge can tell a real TOC field from a typed list."""
    word = (instr.strip().split() or [""])[0].upper()
    if word == "PAGE":
        return "[Page]"
    if word in ("NUMPAGES", "SECTIONPAGES"):
        return "[Page count]"
    if word == "TOC":
        return "[TOC field" + (f', shown text: "{result.strip()}"' if result.strip() else "") + "]"
    return result


class _DocxReader:
    def __init__(self, doc):
        self.doc = doc
        self.numbering = _Numbering(doc)
        self.fields: list[_Field] = []  # fields can span paragraphs (TOC)

    def _emit(self, parts: list, text: str, bold: bool):
        if self.fields:
            if self.fields[-1].in_result:
                self.fields[-1].result.append(text)
            return
        parts.append((text, bold))

    def inline(self, el, parts: list, bold: bool = False, part=None) -> bool:
        """Collect (text, bold) of el's runs into parts; returns True if a page break was found."""
        page_break = False
        part = part or self.doc.part
        for child in el:
            tag = _local(child)
            if tag == "r":
                rpr = child.find(_q("w:rPr"))
                b = rpr.find(_q("w:b")) if rpr is not None else None
                run_bold = b is not None and b.get(_q("w:val"), "true") not in ("0", "false")
                for x in child:
                    t = _local(x)
                    if t == "t":
                        self._emit(parts, x.text or "", run_bold)
                    elif t == "tab":
                        self._emit(parts, " ", False)
                    elif t in ("br", "cr"):
                        if x.get(_q("w:type")) == "page":
                            page_break = True
                        else:
                            self._emit(parts, "\n", False)
                    elif t == "noBreakHyphen":
                        self._emit(parts, "-", run_bold)
                    elif t == "sym":
                        box = x.get(_q("w:char"), "").upper() in ("F0A8", "F06F", "F071")
                        self._emit(parts, "☐" if box else "□", False)
                    elif t == "fldChar":
                        kind = x.get(_q("w:fldCharType"))
                        if kind == "begin":
                            self.fields.append(_Field())
                        elif kind == "separate" and self.fields:
                            self.fields[-1].in_result = True
                        elif kind == "end" and self.fields:
                            f = self.fields.pop()
                            if text := _field_text(f.instr, "".join(f.result)):
                                self._emit(parts, text, False)
                    elif t == "instrText" and self.fields:
                        self.fields[-1].instr += x.text or ""
                    elif t in ("drawing", "pict", "object"):
                        self._emit(parts, self._drawing(x, part), False)
            elif tag == "fldSimple":
                inner: list = []
                self.inline(child, inner, bold, part)
                if text := _field_text(child.get(_q("w:instr"), ""), "".join(t for t, _ in inner)):
                    self._emit(parts, text, False)
            elif tag in ("hyperlink", "ins", "smartTag", "customXml", "sdt", "sdtContent", "dir", "bdo", "moveTo"):
                page_break |= self.inline(child, parts, bold, part)
        return page_break

    def _drawing(self, el, part) -> str:
        for chart in el.iter(_q("c:chart")):
            try:
                xml = etree.fromstring(part.related_parts[chart.get(_q("r:id"))].blob, _PARSER)
                return "\n\n" + chart_text(xml) + "\n\n"
            except Exception:  # noqa: BLE001 - unreadable chart part
                return "[Chart]"
        boxes = list(el.iter(_q("wps:txbx"))) + list(el.iter(_q("v:textbox")))
        if boxes:
            texts = []
            for p in (p for box in boxes for p in box.iter(_q("w:p"))):
                inner: list = []
                self.inline(p, inner, False, part)
                if s := "".join(t for t, _ in inner).strip():
                    texts.append(s)
            return "[Text box: " + " / ".join(texts) + "]"
        return "[Picture]"

    def _style(self, p_el):
        from docx.text.paragraph import Paragraph

        try:
            return Paragraph(p_el, self.doc).style
        except Exception:  # noqa: BLE001 - broken style reference
            return None

    def _num_pr(self, p_el, style) -> tuple[str, int] | None:
        """(numId, level) of the paragraph, else inherited through the style chain (List Number has it there)."""
        ppr = p_el.find(_q("w:pPr"))
        numpr = ppr.find(_q("w:numPr")) if ppr is not None else None
        num_id = ilvl = None
        if numpr is not None:
            n, i = numpr.find(_q("w:numId")), numpr.find(_q("w:ilvl"))
            num_id = n.get(_q("w:val")) if n is not None else None
            ilvl = int(i.get(_q("w:val"))) if i is not None else None
        st, depth = style, 0
        while (num_id is None or ilvl is None) and st is not None and depth < 10:
            sppr = st.element.find(_q("w:pPr"))
            snum = sppr.find(_q("w:numPr")) if sppr is not None else None
            if snum is not None:
                n, i = snum.find(_q("w:numId")), snum.find(_q("w:ilvl"))
                if num_id is None and n is not None:
                    num_id = n.get(_q("w:val"))
                if ilvl is None and i is not None:
                    ilvl = int(i.get(_q("w:val")))
            st, depth = st.base_style, depth + 1
        return None if num_id is None or num_id == "0" else (num_id, ilvl or 0)

    @staticmethod
    def _bold_text(parts: list, bold_allowed: bool) -> str:
        groups: list[list] = []
        for text, bold in parts:
            bold = bold and bold_allowed
            if groups and (groups[-1][1] == bold or not text.strip()):
                groups[-1][0] += text
            else:
                groups.append([text, bold])
        out = []
        for text, bold in groups:
            if bold and text.strip():
                out.append(f"{text[: len(text) - len(text.lstrip())]}**{text.strip()}**{text[len(text.rstrip()):]}")
            else:
                out.append(text)
        return "".join(out)

    def paragraph(self, p_el) -> list[str]:
        style = self._style(p_el)
        name = style.name if style is not None else ""
        parts: list = []
        page_break = self.inline(p_el, parts)
        ppr = p_el.find(_q("w:pPr"))
        heading = name in ("Title", "Subtitle", "TOC Heading") or name.startswith("Heading")
        raw = "".join(t for t, _ in parts)
        text = self._bold_text(parts, not heading).strip()
        out = ["[Page break]"] if ppr is not None and ppr.find(_q("w:pageBreakBefore")) is not None else []
        if raw.strip():
            num = self._num_pr(p_el, style)
            prefix = self.numbering.prefix(*num) if num else None
            level = num[1] if num else 0
            if name.startswith("List") and (m := re.search(r"(\d)$", name)):
                level = max(level, int(m.group(1)) - 1)
            if prefix is None and name.startswith("List Bullet"):
                prefix = "-"
            elif prefix is None and name.startswith("List Number"):
                prefix = "1."
            if name == "Title":
                out.append(f"# {text}")
            elif name.startswith("Heading"):
                n = int(name.split()[-1]) if name.split()[-1].isdigit() else 1
                out.append("#" * min(n + 1, 6) + " " + (f"{prefix} " if prefix else "") + text)
            elif name == "Subtitle":
                out.append(f"*Subtitle:* {text}")
            elif name == "TOC Heading":
                out.append(f"**{text}** *(TOC heading)*")
            elif prefix is not None:
                out.append("  " * level + (f"{prefix} " if prefix else "") + text.replace("\n", " "))
            elif name in ("Quote", "Intense Quote"):
                out.append("> " + text.replace("\n", "\n> "))
            elif name == "Caption":
                out.append(f"*{text}*")
            else:
                out.append(text)
        if page_break:
            out.append("[Page break]")
        return out

    def cell_text(self, tc) -> str:
        texts = []
        for child in tc:
            if _local(child) == "p":
                parts: list = []
                self.inline(child, parts)
                if s := "".join(t for t, _ in parts).strip():
                    texts.append(s)
            elif _local(child) == "tbl":
                texts.append("[Table: " + " / ".join(self.cell_text(c) for c in child.iter(_q("w:tc"))) + "]")
        return "\n".join(texts)

    def table(self, tbl) -> str:
        rows = []
        for tr in tbl.findall(_q("w:tr")):
            row = []
            for tc in tr.findall(_q("w:tc")):
                tcpr = tc.find(_q("w:tcPr"))
                span = tcpr.find(_q("w:gridSpan")) if tcpr is not None else None
                vmerge = tcpr.find(_q("w:vMerge")) if tcpr is not None else None
                continued = vmerge is not None and vmerge.get(_q("w:val"), "continue") == "continue"
                row.append("" if continued else self.cell_text(tc))
                row += [""] * (int(span.get(_q("w:val"))) - 1 if span is not None else 0)
            rows.append(row)
        return md_table(rows)

    def block(self, el) -> list[str]:
        tag = _local(el)
        if tag == "p":
            return self.paragraph(el)
        if tag == "tbl":
            return [self.table(el)]
        if tag in ("sdt", "sdtContent", "customXml"):
            return [line for child in el for line in self.block(child)]
        return []

    def header_footer(self, kind: str) -> list[str]:
        """Distinct header (or footer) texts of all sections, with the first-page and even-page variants marked."""
        texts = []
        for s in self.doc.sections:
            variants = ((f"first_page_{kind}", " (first page)"), (kind, ""), (f"even_page_{kind}", " (even pages)"))
            for attr, note in variants:
                hf = getattr(s, attr)
                if hf.is_linked_to_previous:
                    continue
                lines = []
                for child in hf._element:
                    tag = _local(child)
                    if tag == "p":
                        parts: list = []
                        self.inline(child, parts, part=hf.part)
                        if t := "".join(x for x, _ in parts).strip():
                            lines.append(t)
                    elif tag == "tbl":
                        lines.append(" | ".join(self.cell_text(tc) for tc in child.iter(_q("w:tc"))))
                if (text := " / ".join(lines)) and (text, note) not in texts:
                    texts.append((text, note))
        return [f"{t}{n}" for t, n in texts]


def text_docx(path: Path) -> str:
    from docx import Document

    doc = Document(path)
    reader = _DocxReader(doc)
    out = [f"*Header:* {h}" for h in reader.header_footer("header")]
    for el in doc.element.body.iterchildren():
        out += reader.block(el)
    if reader.fields:  # field never closed: show what was collected
        rest = "".join("".join(f.result) for f in reader.fields)
        out.append(_field_text(reader.fields[0].instr, rest) or "[unclosed field]")
    out += [f"*Footer:* {f}" for f in reader.header_footer("footer")]
    cleaned: list[str] = []
    for line in out:  # collapse repeated page-break markers
        if not (line == "[Page break]" and cleaned and cleaned[-1] == line):
            cleaned.append(line)
    return "\n\n".join(cleaned)


# ----------------------------------------------------------------------------------------------------------- pptx

def _shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes)
        else:
            yield shape


def _placeholder_kind(shape) -> str | None:
    from pptx.enum.shapes import PP_PLACEHOLDER as P

    if not shape.is_placeholder:
        return None
    try:
        kind = shape.placeholder_format.type
    except Exception:  # noqa: BLE001
        return None
    if kind in (P.TITLE, P.CENTER_TITLE, P.VERTICAL_TITLE):
        return "title"
    if kind == P.SUBTITLE:
        return "subtitle"
    return "footer" if kind in (P.FOOTER, P.SLIDE_NUMBER, P.DATE) else "text"


def text_pptx(path: Path) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(path)
    width, height = prs.slide_width or 0, prs.slide_height or 0
    out = []
    for n, slide in enumerate(prs.slides, start=1):
        head = f"## Slide {n}"
        if layout := (slide.slide_layout.name if slide.slide_layout is not None else ""):
            head += f' (layout "{layout}")'
        if slide._element.get("show") == "0":
            head += " – hidden"
        out.append(head)
        # reading order: top to bottom in bands of 50000 EMU (about 1.4 mm), then left to right
        for shape in sorted(_shapes(slide.shapes), key=lambda s: ((s.top or 0) // 50000, s.left or 0)):
            kind = _placeholder_kind(shape)
            edge = ""
            if width and None not in (shape.left, shape.top, shape.width, shape.height) and (
                    shape.left < -12700 or shape.top < -12700 or shape.left + shape.width > width * 1.01
                    or shape.top + shape.height > height * 1.01):
                edge = " *(reaches over the slide edge)*"
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines = []
                for p in shape.text_frame.paragraphs:
                    if not (text := "".join(r.text for r in p.runs).strip() or p.text.strip()):
                        continue
                    if kind not in ("title", "subtitle", "footer") and (
                            kind == "text" or len(shape.text_frame.paragraphs) > 1 or p.level):
                        text = "  " * p.level + "- " + text
                    lines.append(text)
                if kind == "title":
                    out.append("**Title:** " + " / ".join(lines) + edge)
                elif kind == "subtitle":
                    out.append("**Subtitle:** " + " / ".join(lines) + edge)
                elif kind == "footer":
                    out.append("*Footer area:* " + " / ".join(lines))
                else:
                    out.append("\n".join(lines) + edge)
            if getattr(shape, "has_table", False) and shape.has_table:
                out.append(md_table([[c.text for c in row.cells] for row in shape.table.rows]) + edge)
            if getattr(shape, "has_chart", False) and shape.has_chart:
                out.append(chart_text(shape.chart._chartSpace, where=edge))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                out.append("[Picture]" + edge)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            if notes := slide.notes_slide.notes_text_frame.text.strip():
                out.append("*Notes:* " + notes.replace("\n", " / "))
    return "\n\n".join(out)


# ----------------------------------------------------------------------------------------------------------- xlsx

NUMBER_AS_TEXT = re.compile(r"^\s*[-+]?(\d{1,3}([.\s]\d{3})+|\d+)([.,]\d+)?\s*(€|eur|%)?\s*$", re.I)
FORMULA_AS_TEXT = re.compile(r"^\s*[\"'„“]?\s*=")
REF = re.compile(r"^(?:(?P<sheet>'(?:[^']|'')+'|[^!]+)!)?(?P<range>\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)$")


def _rels(z: zipfile.ZipFile, part: str) -> dict[str, str]:
    """Relationships of a package part: {rId: part name}."""
    try:
        xml = etree.fromstring(z.read(f"{posixpath.dirname(part)}/_rels/{posixpath.basename(part)}.rels".lstrip("/")),
                               _PARSER)
    except KeyError:
        return {}
    out = {}
    for rel in xml:
        target = rel.get("Target") or ""
        if rel.get("TargetMode") != "External":
            out[rel.get("Id")] = (target.lstrip("/") if target.startswith("/")
                                  else posixpath.normpath(posixpath.join(posixpath.dirname(part), target)))
    return out


def xlsx_charts(path: Path) -> dict[str, list[tuple[str, object]]]:
    """Charts per sheet name: [(anchor cell, c:chartSpace element)], read straight from the package (openpyxl does
    not load charts)."""
    out: dict[str, list] = {}
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        try:
            wb = etree.fromstring(z.read("xl/workbook.xml"), _PARSER)
        except KeyError:
            return out
        wb_rels = _rels(z, "xl/workbook.xml")
        sheets = wb.find(_q("s:sheets"))
        for sheet in sheets if sheets is not None else []:
            if (part := wb_rels.get(sheet.get(_q("r:id")))) not in names:
                continue
            sheet_rels = _rels(z, part)
            for drawing in etree.fromstring(z.read(part), _PARSER).iter(_q("s:drawing")):
                if (d_part := sheet_rels.get(drawing.get(_q("r:id")))) not in names:
                    continue
                d_rels = _rels(z, d_part)
                for anchor in etree.fromstring(z.read(d_part), _PARSER):
                    start = anchor.find(_q("xdr:from"))
                    pos = "" if start is None else (f"{get_column_letter(int(start.findtext(_q('xdr:col')) or 0) + 1)}"
                                                    f"{int(start.findtext(_q('xdr:row')) or 0) + 1}")
                    for chart in anchor.iter(_q("c:chart")):
                        if (c_part := d_rels.get(chart.get(_q("r:id")))) in names:
                            space = etree.fromstring(z.read(c_part), _PARSER)
                            out.setdefault(sheet.get("name"), []).append((pos, space))
    return out


def _ranges(cells: list[tuple[int, int]]) -> str:
    """Compact cell list: contiguous runs per column ("B2:B7, D2:D7"), at most 10."""
    columns: dict[int, list[int]] = {}
    for r, c in cells:
        columns.setdefault(c, []).append(r)
    parts = []
    for c in sorted(columns):
        rows, col = sorted(columns[c]), get_column_letter(c)
        start = prev = rows[0]
        for r in rows[1:] + [None]:
            if r is not None and r == prev + 1:
                prev = r
                continue
            parts.append(f"{col}{start}" if start == prev else f"{col}{start}:{col}{prev}")
            start = prev = r
    return ", ".join(parts[:10]) + (f", … (+{len(parts) - 10})" if len(parts) > 10 else "")


def _cell_value(cell, values_sheet) -> str:
    """Formula with its recalculated result ("=B2-C2 → 450"); text that a reader would take for a number or a formula
    is marked, since that is exactly the defect the image hides."""
    from openpyxl.worksheet.formula import ArrayFormula

    v = cell.value
    if isinstance(v, ArrayFormula):
        v = v.text
    if cell.data_type == "f":
        if values_sheet is None:
            return str(v)
        result = values_sheet[cell.coordinate].value
        return f"{v} → {_value(result) if result is not None else '(empty)'}"
    if v is None:
        return ""
    if isinstance(v, str):
        if FORMULA_AS_TEXT.match(v):
            return f"`{v}` (text, not a formula)"
        return f"`{v}` (text)" if NUMBER_AS_TEXT.match(v) else v
    return _value(v)


def _resolver(wb_values, wb_formulas, default_sheet: str):
    """resolve(formula) for chart_text: the cells of a chart range, recalculated if possible."""
    def resolve(formula: str):
        m = REF.match((formula or "").strip())
        if not m:
            return None
        sheet = m.group("sheet") or default_sheet
        if sheet.startswith("'"):
            sheet = sheet[1:-1].replace("''", "'")
        source = wb_values if wb_values is not None else wb_formulas
        if sheet not in source.sheetnames:
            return None
        cells = source[sheet][m.group("range").replace("$", "")]
        if not isinstance(cells, tuple):
            cells = [cells]
        else:
            cells = [c for row in cells for c in (row if isinstance(row, tuple) else (row,))]
        return [f"{c.value} (not calculated)" if wb_values is None and c.data_type == "f" else c.value for c in cells]

    return resolve


def text_xlsx(path: Path, recalculated: Path | None = None) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    values = load_workbook(recalculated, data_only=True) if recalculated and Path(recalculated).exists() else None
    try:
        charts = xlsx_charts(path)
    except (zipfile.BadZipFile, etree.XMLSyntaxError):
        charts = {}
    out = []
    if values is None and any(c.data_type == "f" for ws in wb.worksheets for row in ws.iter_rows() for c in row):
        out.append("*(Formulas without results: file not recalculated)*")
    for ws in wb.worksheets:
        out.append(f'## Sheet "{ws.title}"' + (" (hidden)" if ws.sheet_state != "visible" else ""))
        values_sheet = values[ws.title] if values is not None and ws.title in values.sheetnames else None
        used = [(c.row, c.column) for row in ws.iter_rows() for c in row if c.value is not None]
        if used:
            r0, r1 = min(r for r, _ in used), max(r for r, _ in used)
            c0, c1 = min(c for _, c in used), max(c for _, c in used)
            rows = [[""] + [get_column_letter(c) for c in range(c0, c1 + 1)]]
            rows += [[str(r)] + [_cell_value(ws.cell(r, c), values_sheet) for c in range(c0, c1 + 1)]
                     for r in range(r0, r1 + 1)]
            out.append(md_table(rows))
        else:
            out.append("(empty sheet)")
        comments, formats = [], {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and c.number_format != "General" and c.data_type in ("n", "f", "d"):
                    formats.setdefault(c.number_format, []).append((c.row, c.column))
                if c.comment is not None and c.comment.text.strip():
                    comments.append(f"{c.coordinate}: {c.comment.text.strip()}")
        if formats:
            out.append("Number formats: " + "; ".join(f"`{f}` {_ranges(k)}" for f, k in formats.items()))
        if ws.merged_cells.ranges:
            out.append("Merged cells: " + ", ".join(str(m) for m in ws.merged_cells.ranges))
        for label, hf in (("Header", ws.oddHeader), ("Footer", ws.oddFooter)):
            if parts := [t for t in (hf.left.text, hf.center.text, hf.right.text) if t]:
                out.append(f"*{label} (print):* " + " | ".join(parts))
        if comments:
            out.append("Comments: " + "; ".join(comments))
        resolve = _resolver(values, wb, ws.title)
        for anchor, space in charts.get(ws.title, []):
            out.append(chart_text(space, resolve, where=f" (at {anchor})" if anchor else ""))
    return "\n\n".join(out)


# ---------------------------------------------------------------------------------------------------------- text API

def render_text(path, recalculated: Path | None = None) -> str:
    """Markdown text of a docx/xlsx/pptx. Needs no LibreOffice: for xlsx, `recalculated` is a LibreOffice-recalculated
    copy (checks.xlsx.recalculate) that supplies the formula results; without it formulas show without results.
    For many files use render_texts, which recalculates all xlsx with one LibreOffice start."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return text_docx(path)
    if suffix == ".pptx":
        return text_pptx(path)
    if suffix == ".xlsx":
        return text_xlsx(path, recalculated)
    raise ValueError(f"unknown file type: {path.name}")


def render_texts(paths, recalculate_xlsx: bool = True) -> dict[Path, str]:
    """{path: text} for many files; all xlsx are recalculated with one LibreOffice run (if recalculate_xlsx). A file
    that cannot be read gets an "[File not readable: …]" text instead of stopping the batch: for the judge that is a
    finding, not a crash."""
    paths = [Path(p) for p in dict.fromkeys(paths)]
    with tempfile.TemporaryDirectory(prefix="render-") as tmp:
        xlsx = [p for p in paths if p.suffix.lower() == ".xlsx"]
        recalculated = recalculate(xlsx, Path(tmp)) if recalculate_xlsx and xlsx else {}
        out = {}
        for p in paths:
            try:
                out[p] = render_text(p, recalculated.get(p))
            except Exception as e:  # noqa: BLE001 - a broken file must not stop the batch
                out[p] = f"[File not readable: {type(e).__name__}: {e}]"
        return out


# -------------------------------------------------------------------------------------------------- neutral copies

NEUTRAL_TIME = (2026, 1, 1, 0, 0, 0)
NEUTRAL_ISO = "2026-01-01T00:00:00Z"
OFFICE_SUFFIXES = (".docx", ".xlsx", ".pptx", ".docm", ".xlsm", ".pptm")
_CORE_TIME = re.compile(r"(<(dcterms:created|dcterms:modified|cp:lastPrinted)\b[^>]*>)[^<]*(</\2>)")


def neutral_bytes(data: bytes) -> bytes:
    """Office file (ZIP) without traces of when it was built: fixed ZIP entry times and core.xml dates.

    The runs of different models ran at different times, so build times would reveal the model to a judge that gets
    the file. Nested Office files (the chart workbooks in ppt/embeddings/, word/embeddings/) are treated the same way.
    """
    with zipfile.ZipFile(io.BytesIO(data)) as src:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as dst:
            for info in src.infolist():
                content, name = src.read(info), info.filename
                if name.lower().endswith("core.xml") and "docprops/" in name.lower():
                    content = _CORE_TIME.sub(lambda m: f"{m.group(1)}{NEUTRAL_ISO}{m.group(3)}",
                                             content.decode("utf-8")).encode("utf-8")
                elif name.lower().endswith(OFFICE_SUFFIXES):
                    try:
                        content = neutral_bytes(content)
                    except zipfile.BadZipFile:
                        pass
                new = zipfile.ZipInfo(name, date_time=NEUTRAL_TIME)
                new.compress_type, new.external_attr = info.compress_type, info.external_attr
                dst.writestr(new, content)
        return buffer.getvalue()


def neutral_copy(path, dst) -> None:
    """Copy of a file for the judge: Office files through neutral_bytes (a broken ZIP is copied as it is, since the
    judge should see the broken file), other files unchanged."""
    data = Path(path).read_bytes()
    if Path(path).suffix.lower() in OFFICE_SUFFIXES:
        try:
            data = neutral_bytes(data)
        except zipfile.BadZipFile:
            pass
    Path(dst).write_bytes(data)


# ---------------------------------------------------------------------------------------------------------- images

def _soffice(args: list[str], folder: Path, timeout: float) -> None:
    """LibreOffice headless with a fresh profile in folder/profile that recalculates xlsx when loading (so the PDF
    shows formula results, which openpyxl files do not carry), via the content check's process-group runner."""
    if not (Path(SOFFICE).exists() or shutil.which(SOFFICE)):
        raise RuntimeError(f"LibreOffice not found ({SOFFICE})")
    profile = _profile(folder / "profile")
    _run_soffice([SOFFICE, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--norestore", "--nologo",
                  "--nodefault", *args], timeout)


def render_images(paths, out_dir, dpi: int = DPI, xlsx_size: int = XLSX_SIZE,
                  timeout: float = 1800) -> dict[Path, list[Path]]:
    """Page images: {path: [png, …]} in page order ([] if LibreOffice made no PDF of the file).

    One LibreOffice run for all docx/pptx and one for all xlsx (Calc needs its own export filter for one page per
    sheet), then pdftoppm per file: at most 8 Word pages, 12 slides, 12 sheets (MAX_PAGES). The files are rendered
    from neutral copies under neutral names, so equal file names in different folders do not collide and nothing
    in out_dir points to the model. The PNGs are named f<n>-<page>.png in out_dir; PDFs and the LibreOffice profile
    live in a temp folder.
    """
    paths = [Path(p) for p in dict.fromkeys(paths)]
    if unknown := [p.name for p in paths if p.suffix.lower() not in FILE_TYPES]:
        raise ValueError(f"unknown file type: {', '.join(unknown)}")
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out: dict[Path, list[Path]] = {}
    with tempfile.TemporaryDirectory(prefix="render-") as tmp:
        tmp = Path(tmp)
        (tmp / "in").mkdir()
        copies = {p: tmp / "in" / f"f{i:05d}{p.suffix.lower()}" for i, p in enumerate(paths)}
        for p, copy in copies.items():
            neutral_copy(p, copy)
        sheets = [str(c) for c in copies.values() if c.suffix == ".xlsx"]
        others = [str(c) for c in copies.values() if c.suffix != ".xlsx"]
        if others:
            _soffice(["--convert-to", "pdf", "--outdir", str(tmp / "pdf"), *others], tmp, timeout)
        if sheets:
            _soffice(["--convert-to", XLSX_PDF, "--outdir", str(tmp / "pdf"), *sheets], tmp, timeout)
        for p, copy in copies.items():
            pdf, prefix = tmp / "pdf" / f"{copy.stem}.pdf", out_dir / copy.stem
            for old in out_dir.glob(f"{copy.stem}-*.png"):
                old.unlink()
            if not pdf.exists():
                out[p] = []
                continue
            size = ["-scale-to", str(xlsx_size)] if copy.suffix == ".xlsx" else ["-r", str(dpi)]
            subprocess.run(["pdftoppm", "-png", *size, "-l", str(MAX_PAGES[copy.suffix]), str(pdf), str(prefix)],
                           check=False, capture_output=True, timeout=600)
            out[p] = sorted(out_dir.glob(f"{copy.stem}-*.png"))
    return out


def combine_pages(images: list[Path], dst: Path, columns: int) -> None:
    """The pages of one file side by side on a grey background (COLUMNS per file type), so the judge gets one image
    per file. Saved with 256 colours: looks the same, about 60 % smaller. Pillow comes with python-pptx."""
    from PIL import Image

    pages = [Image.open(p).convert("RGB") for p in images]
    w, h = max(i.width for i in pages), max(i.height for i in pages)
    rows, columns = (len(pages) + columns - 1) // columns, min(columns, len(pages))
    canvas = Image.new("RGB", (columns * (w + 8) + 8, rows * (h + 8) + 8), (140, 140, 140))
    for n, page in enumerate(pages):
        canvas.paste(page, (8 + (n % columns) * (w + 8), 8 + (n // columns) * (h + 8)))
    canvas.quantize(256).save(dst, optimize=True)
