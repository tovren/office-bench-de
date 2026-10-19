"""xlsx module internals on small workbooks built with openpyxl: percent format, BW holidays, sum_range next to key
columns, running sums vs. total rows, and the LibreOffice handling (process group kill on timeout, soffice discovery,
a real recalculation as a slow test)."""

import os
import subprocess
import time
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from officebench.checks import xlsx as X
from officebench.checks.content import check_one
from officebench.checks.core import Context
from officebench.tasks import Checks, Task, TaskSpec, TurnSpec


def make_task(file_type: str | None = "xlsx", **checks) -> Task:
    """A task with these [checks] (validated like a task file)."""
    spec = TaskSpec(id="test", title="Test", version=1, kind="file" if file_type else "no_file", file_type=file_type,
                    turns=[TurnSpec(text="Test")], checks=Checks(**checks))
    return Task(spec=spec, turns=(), path=Path("."), hash="")


def test_percent_format():
    assert X.percent_format("0.0%") and X.percent_format("#,##0.00%")
    assert not X.percent_format('#,##0.00" %"') and not X.percent_format("0.00\\%") and not X.percent_format(None)


def test_holidays_bw():
    assert X.easter(2026) == date(2026, 4, 5) and X.easter(2027) == date(2027, 3, 28)
    f = X.holidays_bw(2027)
    assert f[date(2027, 3, 26)] == "Karfreitag" and f[date(2027, 5, 6)] == "Christi Himmelfahrt"
    assert f[date(2027, 5, 27)] == "Fronleichnam" and date(2027, 12, 24) not in f and len(f) == 12


def test_sum_range_with_key_columns():
    """A 'Summe' column that adds only the cost columns, with area, share and head count in front of them, is no
    incomplete sum; a cost column left of the range is."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Abteilung", "Fläche (m²)", "Anteil", "Mitarbeitende", "Heizung (€)", "Strom (€)", "Reinigung (€)",
               "Summe"])
    ws.append(["Verwaltung", 220, 0.22, 12, 4048, 1440, 2112, "=SUM(E2:G2)"])
    ws.append(["Vertrieb", 140, 0.14, 8, 2576, 960, 1344, "=SUM(F3:G3)"])  # leaves out Heizung
    ws["C2"].number_format = ws["C3"].number_format = "0%"
    for cell in ("E2", "F2", "G2", "E3", "F3", "G3"):
        ws[cell].number_format = '#,##0 "€"'
    ctx = Context(make_task())
    X.sum_range(X.Sheet(ws), ctx)
    assert len(ctx.defects) == 1 and "H3 =SUM(F3:G3) leaves out E3" in ctx.defects[0] and "H2" not in ctx.defects[0]


def _running_sum(path: Path, with_total_row: bool):
    wb = Workbook()
    ws = wb.active
    ws.append(["Monat", "Einnahmen", "Ausgaben", "Saldo", "Saldo kumuliert"])
    for i, m in enumerate(["Jan", "Feb", "Mär", "Apr", "Mai", "Jun"], 2):
        ws.append([m, 3000, 2500, f"=B{i}-C{i}", f"=SUM(D$2:D{i})"])
    if with_total_row:
        ws.append(["Summe", "=SUM(B2:B7)", "=SUM(C2:C7)", "=SUM(D2:D7)"])
    wb.save(path)


def test_running_sum_is_no_total_row(tmp_path):
    without, with_ = tmp_path / "without.xlsx", tmp_path / "with.xlsx"
    _running_sum(without, False)
    _running_sum(with_, True)
    b = X.Sheet(load_workbook(without).active)
    assert b.data == [2, 3, 4, 5, 6, 7] and b.totals == []
    assert X.Sheet(load_workbook(with_).active).totals == [8]
    t = make_task(total_row=True, month_rows_min=6)
    d_without = check_one(without, Context(t), {}).defects
    d_with = check_one(with_, Context(t), {}).defects
    assert [m.check for m in d_without if m.hard] == ["not_recalculated", "total_row"]
    assert [m.check for m in d_with if m.hard] == ["not_recalculated"]


def test_libreoffice_timeout_kills_process_group(tmp_path):
    """A hanging soffice (script that starts a child and sleeps) is killed with its whole process group."""
    pids = tmp_path / "pids"
    fake = tmp_path / "soffice"
    fake.write_text(f"#!/bin/sh\nsleep 60 &\necho $! >> {pids}\necho $$ >> {pids}\nsleep 60\n")
    fake.chmod(0o755)
    path = tmp_path / "a.xlsx"
    Workbook().save(path)
    start = time.time()
    result = X.recalculate([path], tmp_path / "tmp", soffice=str(fake), timeout=0.5)
    assert result == {path: None} and time.time() - start < 10  # first try and retry, both killed after 0.5 s
    time.sleep(0.2)
    for pid in map(int, pids.read_text().split()):
        try:
            os.kill(pid, 0)
            alive = subprocess.run(["ps", "-o", "stat=", "-p", str(pid)], capture_output=True, text=True).stdout.strip()
            assert alive.startswith("Z") or not alive, f"process {pid} still running"
        except ProcessLookupError:
            pass


def test_find_soffice_falls_back_to_macos_bundle(monkeypatch):
    monkeypatch.setattr(X.shutil, "which", lambda name: None)
    monkeypatch.setattr(X.Path, "exists", lambda self: str(self) == X._MACOS_SOFFICE)
    assert X._find_soffice() == X._MACOS_SOFFICE
    monkeypatch.setattr(X.Path, "exists", lambda self: False)
    assert X._find_soffice() == "/usr/bin/soffice"


@pytest.mark.slow
@pytest.mark.skipif(not Path(X.SOFFICE).exists(), reason="LibreOffice not installed")
def test_recalculate_with_libreoffice(tmp_path):
    """openpyxl stores no formula results; after the recalculation the values are there (also #DIV/0!)."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Menge", "Preis", "Gesamt"])
    ws.append([3, 2.5, "=A2*B2"])
    ws.append([1, 0, "=A3/B3"])
    path = tmp_path / "r.xlsx"
    wb.save(path)
    result = X.recalculate([path], tmp_path / "tmp")
    assert result[path] is not None
    values = load_workbook(result[path], data_only=True).active
    assert values["C2"].value == 7.5 and values["C3"].value == "#DIV/0!"
