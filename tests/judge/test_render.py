"""judge/render.py: Markdown text (header/footer, numbering, charts, notes, formulas), page images, neutral copies.
Realistic inputs are built from the reference solutions with the real tool; small hand-built files cover the cases
the references do not contain (TOC field, text that looks like a number or formula, hidden layout details)."""

import io
import json
import zipfile
from datetime import datetime
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from officebench.checks.xlsx import SOFFICE
from officebench.judge import render as R
from officebench.tasks import load_tasks
from officebench.tool import build

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
TASKS = load_tasks()
REFERENCES = sorted(i for i, t in TASKS.items() if t.reference_path.exists())
needs_libreoffice = pytest.mark.skipif(not Path(SOFFICE).exists(), reason="LibreOffice not installed")


def reference_file(task_id: str, folder: Path, name: str | None = None) -> Path:
    """The file built from tasks/<id>/reference.json, saved as <folder>/<name or id>.<suffix>."""
    call = json.loads(TASKS[task_id].reference_path.read_text(encoding="utf-8"))
    built = build(call["function"], call["arguments"]).files[0]
    path = folder / f"{name or task_id}{Path(built.name).suffix}"
    path.write_bytes(built.data)
    return path


# ------------------------------------------------------------------------------------------------------------ text

def test_docx(tmp_path):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Personalabteilung"
    footer = doc.sections[0].footer.paragraphs[0]
    footer.add_run("Seite ")
    footer._p.append(parse_xml(f'<w:fldSimple {W} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))
    doc.add_heading("Protokoll", level=0)
    doc.add_heading("Beschlüsse", level=1)
    for text in ("Erster Schritt", "Zweiter Schritt"):
        doc.add_paragraph(text, style="List Number")
    doc.add_paragraph("Punkt", style="List Bullet")
    p = doc.add_paragraph("Frist: ")
    p.add_run("Freitag").bold = True
    p.add_run().add_break(WD_BREAK.PAGE)
    table = doc.add_table(rows=2, cols=2)
    for (r, c), text in {(0, 0): "Wer", (0, 1): "Was", (1, 0): "Anna", (1, 1): "Liste | Plan"}.items():
        table.cell(r, c).text = text
    path = tmp_path / "t.docx"
    doc.save(path)
    text = R.render_text(path)
    lines = text.split("\n\n")
    assert lines[0] == "*Header:* Personalabteilung"
    assert lines[-1] == "*Footer:* Seite [Page]"
    assert "# Protokoll" in lines and "## Beschlüsse" in lines
    assert "1. Erster Schritt" in lines and "2. Zweiter Schritt" in lines and "- Punkt" in lines
    assert "Frist: **Freitag**" in lines and "[Page break]" in lines
    assert "| Anna | Liste \\| Plan |" in text


def test_docx_toc_field(tmp_path):
    doc = Document()
    p = doc.add_paragraph()._p
    for xml in (f'<w:r {W}><w:fldChar w:fldCharType="begin"/></w:r>',
                f'<w:r {W}><w:instrText> TOC \\o "1-3" </w:instrText></w:r>',
                f'<w:r {W}><w:fldChar w:fldCharType="separate"/></w:r>',
                f'<w:r {W}><w:t>Felder aktualisieren</w:t></w:r>', f'<w:r {W}><w:fldChar w:fldCharType="end"/></w:r>'):
        p.append(parse_xml(xml))
    path = tmp_path / "toc.docx"
    doc.save(path)
    assert R.render_text(path).strip() == '[TOC field, shown text: "Felder aktualisieren"]'


def test_xlsx(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget 2026"
    ws.append(["Monat", "Einnahmen", "Ausgaben", "Saldo"])
    for i, (m, e, a) in enumerate((("Januar", 3000, 2500), ("Februar", 3100, 2600)), start=2):
        ws.append([m, e, a, f"=B{i}-C{i}"])
        ws[f"B{i}"].number_format = '#,##0 "€"'
    ws["A4"] = "Summe"
    ws["B4"] = '"=SUM(B2:B3)"'  # formula stored as text (the Gemma bug)
    ws["C4"] = "5.100"  # number stored as text
    chart = BarChart()
    chart.type = "col"
    chart.title = "Einnahmen und Ausgaben"
    chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=3))
    ws.add_chart(chart, "F2")
    ws.merge_cells("A6:D6")
    path = tmp_path / "t.xlsx"
    wb.save(path)
    text = R.render_text(path)
    assert '## Sheet "Budget 2026"' in text
    assert "|  | A | B | C | D |" in text and "| 2 | Januar | 3000 | 2500 | =B2-C2 |" in text
    assert '`"=SUM(B2:B3)"` (text, not a formula)' in text and "`5.100` (text)" in text
    assert 'Number formats: `#,##0 "€"` B2:B3' in text and "Merged cells: A6:D6" in text
    assert '**Chart** (at F2): column chart (barChart, col); title "Einnahmen und Ausgaben"' in text
    assert "Series \"Einnahmen\" ['Budget 2026'!$B$2:$B$3]: 3000; 3100" in text
    assert "Categories ['Budget 2026'!$A$2:$A$3]: Januar; Februar" in text
    assert "Formulas without results" in text
    # a recalculated copy supplies the results
    recalculated = tmp_path / "recalculated.xlsx"
    wb2 = load_workbook(path)
    wb2["Budget 2026"]["D2"] = wb2["Budget 2026"]["D3"] = 500
    wb2.save(recalculated)
    with_results = R.render_text(path, recalculated)
    assert "=B2-C2 → 500" in with_results and "Formulas without results" not in with_results


def test_pptx(tmp_path):
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "Vertrieb Q3"
    title.placeholders[1].text = "Zahlen und Maßnahmen"
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Maßnahmen"
    tf = slide.placeholders[1].text_frame
    tf.text = "Neukunden"
    p = tf.add_paragraph()
    p.text = "zwei Messen"
    p.level = 1
    slide.notes_slide.notes_text_frame.text = "Sprechernotiz"
    chart_slide = prs.slides.add_slide(prs.slide_layouts[5])
    chart_slide.shapes.title.text = "Umsatz"
    data = CategoryChartData()
    data.categories = ["Nord", "Süd"]
    data.add_series("Q3 2026", (1.5, 2.0))
    data.add_series("Q3 2025", (1.2, 1.8))
    chart_slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(1), Inches(2), Inches(12), Inches(4), data)
    path = tmp_path / "t.pptx"
    prs.save(path)
    text = R.render_text(path)
    assert '## Slide 1 (layout "Title Slide")' in text
    assert "**Title:** Vertrieb Q3" in text and "**Subtitle:** Zahlen und Maßnahmen" in text
    assert "- Neukunden\n  - zwei Messen" in text and "*Notes:* Sprechernotiz" in text
    assert "bar chart (barChart, bar)" in text and "*(reaches over the slide edge)*" in text
    assert "Categories [Sheet1!$A$2:$A$3]: Nord; Süd" in text
    assert 'Series "Q3 2026" [Sheet1!$B$2:$B$3]: 1.5; 2' in text and 'Series "Q3 2025"' in text


def test_unknown_type(tmp_path):
    with pytest.raises(ValueError):
        R.render_text(tmp_path / "x.odt")
    with pytest.raises(ValueError):
        R.render_images([tmp_path / "x.odt"], tmp_path)


@pytest.mark.parametrize("task_id", REFERENCES)
def test_reference_files(task_id, tmp_path):
    """Every reference file renders; the structure matches the file type and nothing is lost to an exception."""
    path = reference_file(task_id, tmp_path)
    text = R.render_text(path)
    heading = {".docx": "", ".xlsx": "## Sheet ", ".pptx": "## Slide 1"}[path.suffix]
    assert text.strip() and heading in text and "not readable" not in text


def test_reference_details(tmp_path):
    """German content stays as written; Word page fields, slide notes and charts come through."""
    mahnung = R.render_text(reference_file("mahnung", tmp_path))
    assert "*Footer:* Zahlungserinnerung Seite [Page] von [Page count]" in mahnung
    assert "| 2026-0412 | 03.08.2026 | 02.09.2026 | 2.380,00 € |" in mahnung and "**7.735,50 €**" in mahnung
    vertrieb = R.render_text(reference_file("vertrieb", tmp_path))
    assert "**Chart**: column chart (barChart, col); no title" in vertrieb
    assert 'Series "2026" [Sheet1!$C$2:$C$6]: 576; 506; 618; 432; 663' in vertrieb
    assert '- Axis titles: value axis "Tsd. €"' in vertrieb and "*Notes:* Insgesamt ein gutes Jahr" in vertrieb


def test_render_texts_without_recalculation(tmp_path):
    """A broken file becomes a finding in its text, the other files still render."""
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"kein zip")
    xlsx = reference_file("filialumsatz", tmp_path)
    texts = R.render_texts([broken, xlsx], recalculate_xlsx=False)
    assert texts[broken].startswith("[File not readable: ")
    assert "=SUM(B2:E2) |" in texts[xlsx] and "Formulas without results" in texts[xlsx]


@pytest.mark.slow
@needs_libreoffice
def test_render_texts_recalculates(tmp_path):
    paths = [reference_file("filialumsatz", tmp_path), reference_file("mahnung", tmp_path)]
    texts = R.render_texts(paths)
    assert "| 2 | Stuttgart | 412380 | 398120 | 441900 | 476250 | =SUM(B2:E2) → 1728650 |" in texts[paths[0]]
    assert "Series \"Jahressumme\" ['Umsätze 2026'!$F$2:$F$9]: 1728650; 1056560" in texts[paths[0]]
    assert "Formulas without results" not in texts[paths[0]] and "# Zahlungserinnerung" in texts[paths[1]]


# ---------------------------------------------------------------------------------------------------------- images

@pytest.mark.slow
@needs_libreoffice
def test_render_images(tmp_path):
    """One page image per Word page / slide / sheet, at the old sizes; equal names in different folders do not
    collide; nothing in out_dir carries the original file names."""
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    docx = reference_file("mahnung", tmp_path / "a", "file")
    pptx = reference_file("vertrieb", tmp_path / "b", "file")
    xlsx = reference_file("auftragsdaten", tmp_path / "a", "file")
    other_docx = reference_file("handbuch", tmp_path / "b", "file")
    out = R.render_images([docx, pptx, xlsx, other_docx], tmp_path / "png")
    assert [len(out[p]) for p in (docx, pptx, xlsx)] == [1, 6, 3]
    assert len(out[other_docx]) >= 2 and not set(out[docx]) & set(out[other_docx])
    assert Image.open(out[docx][0]).size == (373, 527)  # A4 at 45 dpi
    assert Image.open(out[pptx][0]).size == (600, 338)  # 16:9 at 45 dpi
    assert all(max(Image.open(p).size) == R.XLSX_SIZE for p in out[xlsx])
    assert {p.parent for ps in out.values() for p in ps} == {tmp_path / "png"}
    assert not [p for p in (tmp_path / "png").iterdir() if "file" in p.name or p.suffix != ".png"]


@pytest.mark.slow
@needs_libreoffice
def test_render_images_page_limits(tmp_path):
    """At most 8 Word pages and 12 slides; a second render into the same folder replaces the old pages."""
    doc = Document()
    for n in range(12):
        doc.add_paragraph(f"Seite {n + 1}").add_run().add_break(WD_BREAK.PAGE)
    doc.save(tmp_path / "long.docx")
    prs = Presentation()
    for n in range(14):
        prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = f"Folie {n + 1}"
    prs.save(tmp_path / "long.pptx")
    paths = [tmp_path / "long.docx", tmp_path / "long.pptx"]
    out = R.render_images(paths, tmp_path / "png")
    assert [len(out[p]) for p in paths] == [8, 12]
    short = tmp_path / "short.docx"
    Document().save(short)
    assert len(R.render_images([short], tmp_path / "png")[short]) == 1
    assert len(list((tmp_path / "png").glob("f00000-*.png"))) == 1


# -------------------------------------------------------------------------------------------------- neutral copies

def test_neutral_copy(tmp_path):
    """Copies for the judge carry no build time: ZIP entry times and core.xml dates fixed, also in embedded
    workbooks."""
    doc = Document()
    doc.core_properties.created = datetime(2026, 10, 13, 21, 7, 5)
    doc.core_properties.modified = datetime(2026, 10, 13, 21, 7, 6)
    doc.add_paragraph("Hallo")
    doc.save(tmp_path / "a.docx")
    pptx = reference_file("vertrieb", tmp_path, "b")  # has a chart, so an embedded workbook
    xlsx = reference_file("filialumsatz", tmp_path, "c")
    for name in ("a.docx", pptx.name, xlsx.name):
        R.neutral_copy(tmp_path / name, tmp_path / f"neutral-{name}")
        with zipfile.ZipFile(tmp_path / f"neutral-{name}") as z:
            assert {i.date_time for i in z.infolist()} == {R.NEUTRAL_TIME}
            embedded = [e for e in z.namelist() if e.endswith(".xlsx")]
            assert bool(embedded) == (name == pptx.name)
            for entry in z.namelist():
                if entry.endswith("core.xml"):
                    core = z.read(entry).decode()
                    assert "2026-10-13" not in core and R.NEUTRAL_ISO in core
            for entry in embedded:
                with zipfile.ZipFile(io.BytesIO(z.read(entry))) as inner:
                    assert {i.date_time for i in inner.infolist()} == {R.NEUTRAL_TIME}
                    if "docProps/core.xml" in inner.namelist():
                        assert R.NEUTRAL_ISO in inner.read("docProps/core.xml").decode()
    assert Document(tmp_path / "neutral-a.docx").paragraphs[0].text == "Hallo"
    assert len(Presentation(tmp_path / "neutral-b.pptx").slides) == 6
    assert R.render_text(tmp_path / "neutral-c.xlsx") == R.render_text(xlsx)
    (tmp_path / "x.txt").write_text("kein office")
    R.neutral_copy(tmp_path / "x.txt", tmp_path / "y.txt")
    assert (tmp_path / "y.txt").read_text() == "kein office"
    (tmp_path / "broken.docx").write_bytes(b"kein zip")
    R.neutral_copy(tmp_path / "broken.docx", tmp_path / "broken-copy.docx")
    assert (tmp_path / "broken-copy.docx").read_bytes() == b"kein zip"
