"""Findings of the independent review of the content check (15.10.2026): ten good files built with the real tool
(tests/fixtures/content/review_cases.json, cases s1–s10) that earlier versions marked 'NOT ok', and the robustness
and validation fixes. The built files need LibreOffice (slow); the rest is fast. The fixes on short texts (dates,
weekend, literal \\n, placeholders, words_min, relaxed patterns, LibreOffice timeout) are in tests/checks/."""

import json
import re
from datetime import date, datetime

import pytest
from docx import Document as DocxDocument
from helpers import CONTENT, build_call, make_task, only
from openpyxl import Workbook
from pydantic import ValidationError

from officebench.checks import content
from officebench.checks.content import check_file, check_files, check_one, missing_topics
from officebench.checks.core import CRASHED, Context, Document
from officebench.checks.dates import check_dates
from officebench.tasks import Checks, load_tasks

CASES = json.loads((CONTENT / "review_cases.json").read_text(encoding="utf-8"))


# ------------------------------------------------------------------------------ s1–s10 with the tool --

@pytest.fixture(scope="module")
def cases(tmp_path_factory) -> dict:
    folder = tmp_path_factory.mktemp("review")
    entries = {}
    for case in CASES:
        result, path = build_call(case, folder / case["name"])
        assert path is not None, (case["name"], result.content[:500])
        entries[case["name"]] = (path, case["task"])
    found = check_files(list(entries.values()))
    return {name: found[p] for name, (p, _) in entries.items()}


@pytest.mark.slow
@pytest.mark.parametrize("name", [c["name"] for c in CASES])
def test_good_file_without_hard_defects(cases, name):
    f = cases[name]
    assert f.ok, [f"[{m.severity}] {m.check}: {m}" for m in f.defects]


@pytest.mark.slow
@pytest.mark.parametrize("name, check", [
    ("s1_budget_cumulative", "month_rows_min"),  # running sum =SUM(D$2:Dn) made every row a total row
    ("s1_budget_cumulative", "total_row"),
    ("s2_vergleich_notebook", "value_ranges"),  # 'note' matched 'Notebook'
    ("s3_reisekosten_hotel_zero", "zero_sum"),  # a column of zeros adds up to 0
    ("s4_reisekosten_purpose_date", "dates_between"),  # a date in the text of 'Anlass' is no trip date
    ("s7_projektplan_min_max", "calculations"),  # MIN/SUM/MAX total row is no Start + Dauer = Ende row
    ("s8_projektplan_progress_sum", "zero_sum"),
    ("s9_anweisung_unc_path", "literal_newline"),  # \\nas01\backup is a path
    ("s10_einladung_time", "weekday"),  # '11.12. 17:00' is no 11.12.2017
    ("s10_einladung_time", "past_dates"),
    ("s10_einladung_time", "topics"),  # 'meldet euch … an'
])
def test_no_more_false_alarm(cases, name, check):
    assert not only(cases[name].defects, check), [str(m) for m in cases[name].defects]


@pytest.mark.slow
def test_validity_date_is_no_purchase_date(cases):
    # 'Garantie bis' 2027 lies outside 2015–13.10.2026 but is no purchase date: neither hard nor soft
    assert not only(cases["s5_inventar_warranty_until"].defects, "dates_between")


# ------------------------------------------------------------------------------ xlsx structure, fast --

def test_dates_between_hard_only_in_date_columns(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Datum", "Anlass", "Erledigt am", "Garantie bis"])
    ws.append([datetime(2026, 10, 5), "Vorbereitung Kundentag am 05.11.2026", datetime(2026, 11, 3),
               datetime(2028, 1, 1)])
    ws.append([datetime(2026, 11, 2), "Messe", datetime(2026, 10, 9), datetime(2028, 1, 1)])
    path = tmp_path / "d.xlsx"
    wb.save(path)
    ctx = Context(make_task("xlsx", dates_between=["2026-10-01", "2026-10-31"]))
    check_one(path, ctx, {})
    m = only(ctx.defects, "dates_between")
    assert [(x.severity, "A3" in x, "C2" in x) for x in m] == [("hard", True, False), ("soft", False, True)]
    assert not any("05.11.2026" in x or "D2" in x for x in m)  # text and validity column are not counted


def test_empty_template_only_zeros(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Monat", "Einnahmen", "Ausgaben", "Saldo"])
    for i, month in enumerate(["Jan", "Feb", "Mär"], 2):
        ws.append([month, 0, 0, f"=B{i}-C{i}"])
    wb.save(tmp_path / "zero.xlsx")
    ws["B2"] = 3000
    wb.save(tmp_path / "one.xlsx")
    zero = check_one(tmp_path / "zero.xlsx", Context(), {}).defects
    assert [(m.label, m.severity) for m in zero if m.check == "empty_number_column"] == [("only zeros", "hard")]
    assert not only(check_one(tmp_path / "one.xlsx", Context(), {}).defects, "empty_number_column")


# ------------------------------------------------------------------------------------- text, fast --

def test_week_range_in_heading_is_no_calendar_week():
    # the rest of the old test_woche_im_kopf is in tests/checks/test_dates.py (test_week_in_heading)
    ctx = Context()
    check_dates(Document([("title", "Protokoll der Teambesprechung vom 12.10.2026"),
                          ("paragraph 1", "Woche 1-3: Analyse")]), ctx)
    assert not only(ctx.defects, "calendar_week")


def test_manual_table_of_contents_is_soft(tmp_path):
    d = DocxDocument()
    d.add_paragraph("Inhaltsverzeichnis", style="Heading 2")
    for t in ("1. Ausgangslage", "2. Ziele", "3. Kosten"):
        d.add_paragraph(t, style="List Number")
    d.save(tmp_path / "manual.docx")
    DocxDocument().save(tmp_path / "none.docx")
    task = make_task("docx", toc=True)
    assert [(m.severity, m.check) for m in check_file(tmp_path / "manual.docx", task)] == [("soft", "toc")]
    assert [(m.severity, m.check) for m in check_file(tmp_path / "none.docx", task)] == [("hard", "toc")]


def test_task_patterns_match_no_word_parts():
    tasks = load_tasks()
    assert not any(re.search(v["column"], "Preis pro Notebook", re.I)
                   for v in tasks["vergleich"].spec.checks.given()["value_ranges"])
    konzept = tasks["konzept"].spec.checks.given()
    assert not missing_topics("Zielsetzung und Kernziele", {"Ziele": konzept["topics"]["Ziele"]})
    assert konzept["words_min"] == {"hard": 250, "soft": 500}


# -------------------------------------------------------------------------------- robustness, fast --

def test_crash_of_one_check_is_soft(monkeypatch, tmp_path):
    d = DocxDocument()
    d.add_paragraph("System X am Dienstag, 12.10.2026")
    d.save(tmp_path / "a.docx")

    def broken(doc, ctx):
        raise ZeroDivisionError("test")
    monkeypatch.setattr(content, "check_dates", broken)
    m = check_file(tmp_path / "a.docx")
    assert [(x.label, x.severity) for x in m if x.check == "dates"] == [(CRASHED, "soft")]
    assert only(m, "placeholders")  # the other checks still report


@pytest.mark.parametrize("checks, error", [
    ({"value_ranges": [{"column": "bewertung(note", "min": 1, "max": 5}]}, "invalid regex"),
    ({"dates_between": ["2026-13-01", "2026-12-31"]}, "dates_between"),
    ({"dates_between": ["2026-12-31", "2026-01-01"]}, "ascending"),
    ({"months": [0, 6]}, "months"), ({"months": [6, 1]}, "months"), ({"columns": "datum"}, "list"),
    ({"names": ["Anna", 1]}, "string"), ({"words_min": {"hard": 600, "soft": 500}}, "hard must be"),
    ({"calculations": [{"name": "x", "columns": {"a": "a", "b": "b"}, "formula": "a +", "result": "b"}]},
     "not a valid expression"),
    ({"calculations": [{"name": "x", "columns": {"a": "a", "b": "b"}, "formula": "__import__('os')",
                        "result": "b"}]}, "may only use"),
    ({"calculations": [{"name": "x", "columns": {"a": "a", "b": "b"}, "formula": "a", "result": "b", "sheet": 0}]},
     "sheet"),
    ({"calculations": [{"name": "x", "columns": {"a": "a", "b": "b"}, "formula": "a", "result": "b",
                        "sheet": "("}]}, "invalid regex"),
])
def test_task_validation(checks, error):
    with pytest.raises(ValidationError, match=re.escape(error)):
        Checks(**checks)


def test_validation_lets_valid_through():
    make_task("xlsx", words_min={"hard": 250, "soft": 500}, months=[1, 6], dates_between=["2026-10-01", "2026-10-31"],
              calculations=[{"name": "x", "columns": {"a": "a", "b": "b"}, "formula": "(a - 0.5) * 2", "result": "b",
                             "sheet": "Plan"}])


def test_today_stays_reference_date():
    assert content._date("2027-01-05") == date(2027, 1, 5) and content._date(None) == date(2026, 10, 13)
