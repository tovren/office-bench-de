"""xlsx checks that need the LibreOffice recalculation: calculations (also across a transposed table, with dates and
percentages), sum_range, formula errors, formulas as text, total 0, chart type, chart values from formulas on another
sheet, summaries computed with COUNTIF/SUMIF across sheets (groups). Files come from the real tool
(tests/fixtures/content/*.json) or, for mistakes the current tool no longer lets through, from openpyxl.
All files are checked in one check_files call (one LibreOffice run)."""

from datetime import date
from pathlib import Path

import pytest
from helpers import build_fixture, make_task
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, Series

from officebench.checks.content import check_files

pytestmark = pytest.mark.slow

BALANCE = {"name": "Saldo = Einnahmen − Ausgaben",
           "columns": {"einnahmen": "^(?!.*saldo).*einnahm", "ausgaben": "^(?!.*saldo).*ausgab", "saldo": "saldo"},
           "formula": "einnahmen - ausgaben", "result": "saldo"}
BUDGET = dict(months=[1, 6], month_rows_min=6, chart=True, chart_type="column", chart_series_min=2,
              formulas_min=1, total_row=True, calculations=[BALANCE])
MONTHS = ["Januar", "Februar", "März", "April", "Mai", "Juni"]


def _budget_transposed(path: Path):
    """Months as columns, Saldo row as text '=B2-B3' (as Gemma wrote it), total column =SUM(B4:G4) -> 0."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Haushalt"
    ws.append(["Kategorie", *MONTHS, "Summe (6 Monate)"])
    ws.append(["Einnahmen", 3000, 3000, 3100, 3000, 3000, 3200, "=SUM(B2:G2)"])
    ws.append(["Ausgaben", 2200, 2300, 2100, 2400, 2250, 2500, "=SUM(B3:G3)"])
    ws.append(["Saldo"] + [f"={c}2-{c}3" for c in "BCDEFG"] + ["=SUM(B4:G4)"])
    for c in "BCDEFG":
        ws[f"{c}4"].data_type = "s"  # stored as text, not as formula
    wb.save(path)


def _budget_text(path: Path):
    """Formulas as quoted text, no real formula, a line chart with one series."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Monat", "Einnahmen", "Ausgaben", "Saldo"])
    for i, m in enumerate(MONTHS, 2):
        ws.append([m, 3000, 2500, f'"=B{i}-C{i}"'])
    ws.append(["Summe", '"=SUM(B2:B7)"', '"=SUM(C2:C7)"', '"=SUM(D2:D7)"'])
    chart = LineChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=7), titles_from_data=True)
    ws.add_chart(chart, "F2")
    wb.save(path)


def _sums(path: Path):
    """Total row leaves out the last data row, 'Summe' row computing AVERAGE, a sum over numbers stored as text,
    a division by zero, and a total row without Einnahmen/Ausgaben (only the Saldo total)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Monat", "Einnahmen (€)", "Ausgaben (€)", "Saldo (€)", "Quote"])
    for i, m in enumerate(MONTHS, 2):
        ws.append([m, 3000 + 100 * i, 2500, f"=B{i}-C{i}", f"=D{i}/0" if i == 2 else f"=D{i}/B{i}"])
    ws.append(["Summe", None, None, "=SUM(D2:D6)"])  # D7 missing; Einnahmen/Ausgaben totals missing
    ws2 = wb.create_sheet("Sätze")
    ws2.append(["Rolle", "Stundensatz (€)", "Stunden"])
    ws2.append(["Projektleitung", 120, "10"])
    ws2.append(["Entwicklung", 95, "20"])
    ws2.append(["Summe", "=AVERAGE(B2:B3)", "=SUM(C2:C3)"])
    wb.save(path)


def _empty_results(path: Path):
    """A column formula that leaves a note row empty on purpose (=IF(A6="";"";…)) and an IFERROR that hides a
    failed lookup behind an empty cell."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Bestellung"
    ws.append(["Artikel", "Menge", "Preis", "Gesamt", "Lager"])
    for i in range(2, 5):
        ws.append([f"Artikel {i}", i, 2.5, f'=IF(A{i}="","",B{i}*C{i})', f'=IFERROR(VLOOKUP(A{i},G:H,2,0),"")'])
    ws.append(["Summe", None, None, "=SUM(D2:D4)"])
    ws.append([None, "Hinweis: Preise netto", None, '=IF(A6="","",B6*C6)'])
    wb.save(path)


def _savings_rate(path: Path, as_percent_number: bool):
    wb = Workbook()
    ws = wb.active
    ws.append(["Monat", "Einnahmen", "Ausgaben", "Saldo", "Sparquote"])
    for i, m in enumerate(MONTHS, 2):
        rate = f"=D{i}/B{i}*100" if as_percent_number else f"=D{i}/B{i}"
        ws.append([m, 3200, 2800, f"=B{i}-C{i}", rate])
        ws[f"E{i}"].number_format = "0.0" if as_percent_number else "0.0%"
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=7), titles_from_data=True)
    ws.add_chart(chart, "G2")
    wb.save(path)


OTHER_INCOME, EXPENSES = [150, 0, 200, 0, 100, 350], [1880, 1770, 1920, 1860, 1870, 1960]
RATES = [100 * (2800 + s - a) / (2800 + s) for s, a in zip(OTHER_INCOME, EXPENSES, strict=True)]


def _savings_rate_percent_column(path: Path, percent_format: bool):
    """Like a real file (sparquote, Haiku): 'Sparquote in %' = Saldo / Einnahmen gesamt × 100 next to 'Sonstige
    Einnahmen', repeated on a second sheet by formulas; the line chart shows that sheet, the column chart sheet 1
    (series titles as literals, values only as references)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget 2026"
    ws.append(["Monat", "Gehalt", "Sonstige Einnahmen", "Einnahmen gesamt", "Ausgaben gesamt", "Saldo",
               "Sparquote in %"])
    for i, (m, other, exp) in enumerate(zip(MONTHS, OTHER_INCOME, EXPENSES, strict=True), 2):
        ws.append([m, 2800, other, f"=B{i}+C{i}", exp, f"=D{i}-E{i}", f"=F{i}/D{i}*100"])
        ws[f"G{i}"].number_format = "0.00%" if percent_format else "#,##0.00"
    ws.append(["Summe"] + [f"=SUM({c}2:{c}7)" for c in "BCDEF"])
    columns = BarChart()
    for c in (4, 5):
        columns.series.append(Series(Reference(ws, min_col=c, min_row=2, max_row=7), title=ws.cell(1, c).value))
    ws.add_chart(columns, "I2")
    ws2 = wb.create_sheet("Sparquote")
    ws2.append(["Monat", "Sparquote in %"])
    for i, m in enumerate(MONTHS, 2):
        ws2.append([m, f"='Budget 2026'!G{i}"])
    line = LineChart()
    line.series.append(Series(Reference(ws2, min_col=2, min_row=2, max_row=7), title="Sparquote in %"))
    ws2.add_chart(line, "D2")
    wb.save(path)


# the two calculations of tasks/sparquote/task.toml
SPARQUOTE_TASK = [
    {"name": "Saldo = Einnahmen − Ausgaben", "formula": "einnahmen - ausgaben", "result": "saldo",
     "columns": {"einnahmen": "^(?!.*(saldo|quote)).*einnahm", "ausgaben": "^(?!.*(saldo|quote)).*ausgab",
                 "saldo": "^(?!.*quote).*(saldo|differenz|überschuss)"}},
    {"name": "Sparquote = Saldo / Einnahmen", "formula": "saldo / einnahmen", "result": "quote", "tolerance": 0.005,
     "columns": {"einnahmen": "^(?!.*(saldo|quote)).*einnahm", "saldo": "^(?!.*quote).*(saldo|differenz|überschuss)",
                 "quote": "sparquote"}}]


def _inventory(path: Path, range_end: int):
    """Device list and an overview per type on sheet 2 with COUNTIF/SUMIF over sheet 1 (LibreOffice computes them);
    range_end 5 leaves out the last device (a wrong range). Pie chart of the values."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventar"
    ws.append(["Inv.-Nr.", "Gerätetyp", "Kaufpreis (€)"])
    for i, (kind, price) in enumerate([("Laptop", 1000), ("Monitor", 200), ("Laptop", 1200), ("Monitor", 250),
                                       ("Drucker", 500), ("Monitor", 150)], 1):
        ws.append([f"INV-{i:03d}", kind, price])
    ws.append(["Summe", None, "=SUM(C2:C7)"])
    ws2 = wb.create_sheet("Übersicht")
    ws2.append(["Gerätetyp", "Anzahl", "Gesamtwert (€)"])
    end = range_end + 1
    for i, kind in enumerate(["Laptops", "Monitore", "Drucker"], 2):
        ws2.append([kind, f"=COUNTIF(Inventar!$B$2:$B${end},LEFT(A{i},6)&\"*\")",
                    f"=SUMIF(Inventar!$B$2:$B${end},LEFT(A{i},6)&\"*\",Inventar!$C$2:$C${end})"])
    pie = PieChart()
    pie.add_data(Reference(ws2, min_col=3, min_row=1, max_row=4), titles_from_data=True)
    ws2.add_chart(pie, "E2")
    wb.save(path)


GROUPS = [{"name": "Übersicht je Gerätetyp", "source": {"sheet": 1, "group": "typ", "value": "preis"},
           "target": {"sheet": 2, "group": "typ", "count": "anzahl", "value": "wert"}}]
INVENTORY = dict(groups=GROUPS, chart_type="pie", chart_series=["gesamtwert"], chart_values=[[500, 2200, 600]])


def _shares_and_date_text(path: Path):
    """Percent columns: a share × 100 (flagged), a change of +150 % and a workload of 110 % (fine); a date column
    with a valid date as text and the nonexistent 31.11.2026 as text."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Team"
    ws.append(["Name", "Datum", "Anteil am Budget", "Veränderung zum Vorjahr", "Auslastung", "Anteil verplant",
               "Quote Ist"])
    for row in (("Keller", date(2026, 7, 1), 0.4, 1.5, 1.1, 51.5, 1.4), ("Brandt", "2026-07-02", 0.6, -0.2, 0.9, 0.5,
                0.9), ("Lehmann", "31.11.2026", 0.0, 0.1, 1.0, 0.25, 1.0)):
        ws.append(list(row))
    for cells in ws.iter_rows(min_row=2, min_col=3, max_col=7):
        for c in cells:
            c.number_format = "0.0%"
    ws["B2"].number_format = "DD.MM.YYYY"
    wb.save(path)


SAVINGS_RATE = {"name": "Sparquote = Saldo / Einnahmen",
                "columns": {"einnahmen": "^(?!.*(saldo|quote)).*einnahm", "saldo": "^(?!.*quote).*saldo",
                            "quote": "sparquote"},
                "formula": "saldo / einnahmen", "result": "quote", "tolerance": 0.005}


@pytest.fixture(scope="module")
def findings(tmp_path_factory) -> dict:
    folder = tmp_path_factory.mktemp("xlsx")
    files = {
        "budget_good": (build_fixture("budget_good", folder), make_task("xlsx", **BUDGET)),
        "budget_sign": (build_fixture("budget_sign", folder), make_task("xlsx", **BUDGET)),
        "reisekosten": (build_fixture("reisekosten", folder), make_task("xlsx", calculations=[{
            "name": "Erstattung = km × 0,30 €", "columns": {"km": r"kilometer|\bkm\b", "erst": "erstattung"},
            "formula": "km * 0.30", "result": "erst"}], value_ranges=[{"column": "bewertung", "min": 1, "max": 5}])),
        "projektplan": (build_fixture("projektplan", folder), make_task("xlsx", calculations=[{
            "name": "Ende = Start + Dauer", "columns": {"start": "start", "dauer": "dauer", "ende": r"\bende"},
            "formula": "start + dauer - 0.5", "result": "ende", "tolerance": 0.6}, {
            "name": "Ende = Start + Dauer + 3", "columns": {"start": "start", "dauer": "dauer", "ende": r"\bende"},
            "formula": "start + dauer + 3", "result": "ende"}])),
    }
    path = folder / "empty_results.xlsx"
    _empty_results(path)
    files["empty_results"] = (path, make_task("xlsx"))
    for name, builder in (("budget_transposed", _budget_transposed), ("budget_text", _budget_text), ("sums", _sums)):
        path = folder / f"{name}.xlsx"
        builder(path)
        files[name] = (path, make_task("xlsx", **BUDGET) if name != "sums" else make_task("xlsx",
                                                                                          calculations=[BALANCE]))
    for name, percent in (("rate_share", False), ("rate_percent_number", True)):
        path = folder / f"{name}.xlsx"
        _savings_rate(path, percent)
        files[name] = (path, make_task("xlsx", calculations=[SAVINGS_RATE], chart_type="line"))
    for name, percent_format in (("percent_column", False), ("percent_column_formatted", True)):
        path = folder / f"{name}.xlsx"
        _savings_rate_percent_column(path, percent_format)
        files[name] = (path, make_task("xlsx", calculations=SPARQUOTE_TASK, chart_type="line",
                                       chart_series=["sparquote"], chart_values=[RATES], values_present=[36.27]))
    path = folder / "shares.xlsx"
    _shares_and_date_text(path)
    files["shares"] = (path, make_task("xlsx"))
    for name, end in (("inventory", 6), ("inventory_range", 5)):
        path = folder / f"{name}.xlsx"
        _inventory(path, end)
        files[name] = (path, make_task("xlsx", **INVENTORY))
    result = check_files([(p, t) for p, t in files.values()])
    return {name: result[p] for name, (p, _) in files.items()}


def _d(found, check, severity=None):
    return [m for m in found.defects if m.check == check and (severity is None or m.severity == severity)]


def test_budget_good_clean(findings):
    assert findings["budget_good"].clean, [str(m) for m in findings["budget_good"].defects]


def test_balance_with_wrong_sign(findings):
    f = findings["budget_sign"]
    wrong = _d(f, "calculations", "hard")
    assert len(wrong) == 1 and "Budget!D2 = -800, expected 800" in wrong[0] and "3 more" in wrong[0]
    # the tool's total row has no Saldo total: soft
    assert "Budget!D8" in _d(f, "calculations", "soft")[0]


def test_transposed_table_with_formula_as_text(findings):
    f = findings["budget_transposed"]
    assert "Haushalt!H4 = 0, expected 4.550" in _d(f, "calculations", "hard")[0]
    assert "Haushalt!B4:G4" in _d(f, "formula_as_text", "hard")[0]
    assert {m.check for m in f.hard} == {"calculations", "formula_as_text", "chart"}  # SUM column counts as total


def test_formulas_in_quotes(findings):
    f = findings["budget_text"]
    assert '„"=B2-C2"“' in _d(f, "formula_as_text")[0]
    assert {"formulas_min", "total_row", "chart_type", "chart_series_min"} <= {m.check for m in f.hard}
    assert "found: line" in _d(f, "chart_type")[0]
    assert _d(f, "calculations", "soft")  # no row with values: not checkable


def test_sum_range_and_formula_errors(findings):
    f = findings["sums"]
    assert all(m.severity == "soft" for m in _d(f, "sum_range"))
    ranges = " ".join(_d(f, "sum_range"))  # one defect per sheet
    assert "Budget!D8 =SUM(D2:D6) leaves out D7" in ranges and "=AVERAGE(…) in the row „Summe“" in ranges
    assert "Sätze!C4 = 0, the data rows give 30" in ranges  # numbers as text skipped by SUM
    assert "#DIV/0!" in _d(f, "formula_errors", "hard")[0]
    assert _d(f, "zero_sum", "hard") and _d(f, "number_as_text", "soft")
    assert "no value in Budget!B8, C8" in _d(f, "calculations", "soft")[0]


def test_deliberately_empty_formula(findings):
    empty = _d(findings["empty_results"], "formula_empty", "hard")
    assert len(empty) == 1 and "Bestellung!E2:E4" in empty[0] and "D6" not in empty[0], empty


def test_constant_and_value_range(findings):
    f = findings["reisekosten"]
    assert "Fahrten!E4 = 30, expected 24" in _d(f, "calculations", "hard")[0]
    assert "Fahrten!F5 = 7" in _d(f, "value_ranges", "hard")[0]


def test_date_as_serial_number(findings):
    f = findings["projektplan"]
    found = _d(f, "calculations")
    assert len(found) == 1 and "Ende = Start + Dauer + 3" in found[0], [str(m) for m in f.defects]
    assert "Plan!D2 = 30.10.2026, expected 03.11.2026" in found[0]


def test_percent_as_share(findings):
    assert not _d(findings["rate_share"], "calculations")
    wrong = _d(findings["rate_percent_number"], "calculations", "hard")
    assert wrong and "Sheet!E2 = 12,50, expected 0,125" in wrong[0]
    assert _d(findings["rate_share"], "chart_type", "hard")  # bar chart, line expected


def test_percent_scale_with_percent_in_header(findings):
    """'Sparquote in %' = 36,27 without percent format is right (also picks 'Einnahmen gesamt', not 'Sonstige
    Einnahmen'); the chart values come from the formulas on sheet 2."""
    f = findings["percent_column"]
    assert f.clean, [str(m) for m in f.defects]
    wrong = _d(findings["percent_column_formatted"], "calculations", "hard")  # percent format and × 100: 3.627 %
    assert len(wrong) == 1 and "Budget 2026!G2 = 36,27, expected" in wrong[0]


def test_groups_with_formulas_across_sheets(findings):
    assert findings["inventory"].clean, [str(m) for m in findings["inventory"].defects]
    m = _d(findings["inventory_range"], "groups", "hard")
    assert len(m) == 1 and "Monitore: count 2 (Übersicht!B3), per Inventar 3" in m[0]
    assert "Monitore: value 450 (Übersicht!C3), per Inventar 600" in m[0]
    values = _d(findings["inventory_range"], "chart_values", "hard")
    assert "most similar series „Gesamtwert (€)“: 2.200; 450; 500" in values[0]


def test_share_over_100_and_date_as_text(findings):
    f = findings["shares"]
    share = _d(f, "share_over_100", "soft")
    assert len(share) == 1 and "„Anteil verplant“ Team!F2 " in share[0] and "5.150 %" in share[0], share
    assert "Veränderung" not in share[0] and "Auslastung" not in share[0] and "Budget" not in share[0]
    assert "Quote Ist" not in share[0]
    invalid = _d(f, "invalid_date", "hard")
    assert len(invalid) == 1 and "Team!B4" in invalid[0] and "31.11.2026" in invalid[0], invalid
    as_text = _d(f, "date_as_text", "soft")
    assert len(as_text) == 1 and "Team!B3 " in as_text[0], as_text
